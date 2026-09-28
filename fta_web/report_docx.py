"""
The DOCX analysis report (1.7).

Two steps, so the slow part never runs under the state lock and the data can
be tested without Word:

* :func:`collect_report_data` -- pure. Reads a document (anything with
  ``get_data()``, ``get_metadata()`` and an ``analysis`` attribute: a
  ``WebCore`` or a :class:`DocSnapshot`), runs the analyses the chosen
  sections need through :func:`try_call`, and returns plain dicts.
* :func:`build_report` -- turns those dicts into ``.docx`` bytes with
  python-docx, which is imported inside the function: the package is an
  optional extra (``uv sync --extra report``) and the rest of the app must
  import without it.

Cross-workstream analyses (``cutsets.compute``, ``importance.compute``,
``uncertainty.run``, ``lint.run``, ``engine.summary``) are reached through
:func:`try_call`. A module that does not have the function yet, or raises
``NotImplementedError``, makes its section say "unavailable" instead of
failing the whole report.

Headings and labels are localized here (``en``/``ja``); with ``ja`` the
East-Asian font of every style is set to Meiryo so Word does not fall back to
a font without Japanese glyphs.
"""
from __future__ import annotations

import copy
import datetime as _dt
import importlib
import io
from typing import Any, Callable, Dict, List, Optional, Tuple

try:  # normal package import: ``import fta_web.report_docx``
    from .excel_events import event_row, iter_nodes, node_kind
    from .numfmt import clamp_sig_figs, format_prob
except ImportError:  # fallback: ``fta_web/`` itself is on sys.path
    from excel_events import event_row, iter_nodes, node_kind  # type: ignore[no-redef]
    from numfmt import clamp_sig_figs, format_prob  # type: ignore[no-redef]

#: Shown on the title page. The frozen build has no pyproject to read, so the
#: release number lives here; the integrator bumps it with the release.
TOOL_VERSION = "1.7.0"

SECTIONS = (
    "metadata", "headline", "assumptions", "diagram", "events", "cutsets",
    "importance", "uncertainty", "validation", "traceability",
)
#: Sections that need an FTA analysis; skipped (with a note) in ETA mode.
ANALYSIS_SECTIONS = ("headline", "cutsets", "importance", "uncertainty")
DEFAULT_SECTIONS = tuple(s for s in SECTIONS if s != "uncertainty")

DEFAULT_TOP_CUTSETS = 50
DEFAULT_TOP_IMPORTANCE = 30
MAX_TOP_N = 10000
MAX_REPORT_MC_N = 5000
MAX_DIAGRAM_BYTES = 8 * 1024 * 1024
PNG_SIGNATURE = b"\x89PNG\r\n\x1a\n"
DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


# ---- adapters ------------------------------------------------------------------


def _module(name: str):
    """A sibling module by bare name, through the package when possible."""
    package = __package__ or ""
    if package:
        try:
            return importlib.import_module("%s.%s" % (package, name))
        except ImportError:
            pass
    return importlib.import_module(name)


def try_call(module_name: str, fn_name: str, *args: Any, **kwargs: Any
             ) -> Tuple[Any, Optional[str]]:
    """``(result, None)`` or ``(None, reason)``.

    ``reason`` is ``"unavailable"`` when the module or function is missing or
    raises ``NotImplementedError``/``AttributeError`` (a workstream that has
    not landed yet), else ``"failed: <message>"``.
    """
    try:
        module = _module(module_name)
        fn = getattr(module, fn_name)
    except (ImportError, AttributeError):
        return None, "unavailable"
    try:
        return fn(*args, **kwargs), None
    except (NotImplementedError, AttributeError):
        return None, "unavailable"
    except Exception as exc:  # the section fails, not the report
        return None, "failed: %s" % exc


# ---- document snapshot -----------------------------------------------------------


class DocSnapshot:
    """A detached copy of a document, readable outside the state lock."""

    def __init__(self, tree: Any, analysis: Any, metadata: Dict[str, Any],
                 quant_warnings: Optional[List[Dict[str, Any]]] = None):
        self.fta_data = tree
        self.analysis = analysis
        self.title = metadata.get("title", "")
        self.date = metadata.get("date", "")
        self.mode = metadata.get("mode", "FTA") or "FTA"
        self.quant_warnings = quant_warnings or []

    @classmethod
    def of(cls, core) -> "DocSnapshot":
        """Deep copy of ``core``'s document. Caller holds the lock."""
        return cls(
            copy.deepcopy(core.get_data()),
            copy.deepcopy(getattr(core, "analysis", None) or {}),
            dict(core.get_metadata()),
            copy.deepcopy(getattr(core, "quant_warnings", None) or []),
        )

    def get_data(self):
        return self.fta_data

    def get_metadata(self):
        return {"title": self.title, "date": self.date, "mode": self.mode}


# ---- options ------------------------------------------------------------------------


def _int_in(value: Any, lo: int, hi: int, default: int) -> int:
    if isinstance(value, bool):
        return default
    try:
        number = int(value)
    except (TypeError, ValueError, OverflowError):
        return default
    return max(lo, min(hi, number))


def normalize_options(raw: Optional[Dict[str, Any]]) -> Dict[str, Any]:
    """Defaults applied, unknown sections dropped, numbers clamped."""
    raw = raw or {}
    sections = raw.get("sections")
    if isinstance(sections, (list, tuple)):
        chosen = [s for s in SECTIONS if s in {str(x) for x in sections}]
    else:
        chosen = list(DEFAULT_SECTIONS)
    run_mc = bool(raw.get("runUncertainty"))
    if (run_mc or raw.get("uncertainty")) and "uncertainty" not in chosen and not isinstance(sections, (list, tuple)):
        chosen.append("uncertainty")
    top = raw.get("topN") if isinstance(raw.get("topN"), dict) else {}
    lang = str(raw.get("lang") or "en").lower()
    return {
        "sections": chosen,
        "sigFigs": clamp_sig_figs(raw.get("sigFigs")),
        "lang": "ja" if lang.startswith("ja") else "en",
        "topCutsets": _int_in(top.get("cutsets", raw.get("topCutsets")), 1, MAX_TOP_N, DEFAULT_TOP_CUTSETS),
        "topImportance": _int_in(top.get("importance", raw.get("topImportance")), 1, MAX_TOP_N,
                                 DEFAULT_TOP_IMPORTANCE),
        "runUncertainty": run_mc,
        "uncertaintyN": _int_in(raw.get("uncertaintyN"), 1, MAX_REPORT_MC_N, MAX_REPORT_MC_N),
        "uncertainty": raw.get("uncertainty") if isinstance(raw.get("uncertainty"), dict) else None,
        "limits": raw.get("limits") if isinstance(raw.get("limits"), dict) else None,
        "diagramPng": raw.get("diagramPng") if isinstance(raw.get("diagramPng"), (bytes, bytearray)) else None,
        "generated": raw.get("generated"),
    }


# ---- data collection -------------------------------------------------------------------


def _formulas() -> Dict[str, str]:
    try:
        return dict(_module("engine")._FORMULAS)
    except Exception:
        return {"fixed": "q", "rate": "1 - exp(-λT)", "standby": "min(1, λτ/2)",
                "repairable": "λ/(λ+μ)"}


def _repeated_names(repeated: Any, names: Dict[str, str]) -> List[str]:
    out = []
    for item in repeated or []:
        if isinstance(item, dict):
            nid = str(item.get("id", ""))
            out.append(item.get("name") or names.get(nid, nid))
        else:
            out.append(names.get(str(item), str(item)))
    return out


def collect_report_data(core, session_warnings: Optional[List[Dict[str, Any]]] = None,
                        options: Optional[Dict[str, Any]] = None) -> Dict[str, Any]:
    """Everything the report shows, as plain data. Does not modify ``core``."""
    opts = options if options and options.get("_normalized") else normalize_options(options)
    sections = list(opts["sections"])
    tree = copy.deepcopy(core.get_data())
    analysis = copy.deepcopy(getattr(core, "analysis", None) or {})
    meta = dict(core.get_metadata())
    mode = meta.get("mode") or "FTA"
    eta = mode == "ETA"
    generated = opts.get("generated") or _dt.datetime.now().strftime("%Y-%m-%d %H:%M")

    data: Dict[str, Any] = {
        "meta": {
            "title": meta.get("title") or "",
            "date": meta.get("date") or "",
            "mode": mode,
            "generated": generated,
            "version": TOOL_VERSION,
        },
        "sections": sections,
        "sigFigs": opts["sigFigs"],
        "lang": opts["lang"],
        "analysis": analysis,
        "unavailable": {},
        "skippedEta": [s for s in sections if eta and s in ANALYSIS_SECTIONS],
        "diagramPng": opts.get("diagramPng"),
    }
    names = {str(n.get("id")): str(n.get("name", "")) for n, _p, _d in iter_nodes(tree)}
    wanted = set(sections) - set(data["skippedEta"])

    summary = None
    if not eta and wanted & {"headline", "assumptions"}:
        summary, reason = try_call("engine", "summary", tree, analysis)
        if summary is None:
            data["unavailable"]["headline"] = reason
    data["summary"] = summary

    if "assumptions" in wanted:
        formulas = _formulas()
        models: Dict[str, int] = {}
        for node, _p, _d in iter_nodes(tree):
            if node_kind(node) != "event":
                continue
            quant = node.get("quant") if isinstance(node.get("quant"), dict) else {}
            model = str(quant.get("model") or "fixed").lower()
            models[model] = models.get(model, 0) + 1
        data["assumptions"] = {
            "missionTime": analysis.get("missionTime"),
            "timeUnit": analysis.get("timeUnit") or "h",
            "models": [{"model": m, "count": c, "formula": formulas.get(m, "")}
                       for m, c in sorted(models.items())],
            "approximations": copy.deepcopy((summary or {}).get("approximations") or []),
            "nonCoherent": bool((summary or {}).get("nonCoherent")),
            "repeatedEvents": _repeated_names((summary or {}).get("repeatedEvents"), names),
            "cutsets": copy.deepcopy(analysis.get("cutsets") or {}),
        }

    if "events" in wanted:
        data["events"] = [event_row(n, p, d) for n, p, d in iter_nodes(tree)
                          if node_kind(n) == "event"]

    cut_result = None
    if wanted & {"cutsets", "importance"}:
        cut_result, reason = try_call("cutsets", "compute", tree, analysis, opts.get("limits"))
        if cut_result is None:
            for s in ("cutsets", "importance"):
                if s in wanted:
                    data["unavailable"][s] = reason
    if "cutsets" in wanted and cut_result is not None:
        rows = list(cut_result.get("cutSets") or [])
        data["cutsets"] = {
            "rows": rows[: opts["topCutsets"]],
            "total": cut_result.get("total", len(rows)),
            "shown": min(len(rows), opts["topCutsets"]),
            "truncated": bool(cut_result.get("truncated")),
            "truncatedBy": cut_result.get("truncatedBy"),
            "mcub": cut_result.get("mcub"),
            "rareEvent": cut_result.get("rareEvent"),
        }
    if "importance" in wanted and cut_result is not None:
        measures, reason = try_call("importance", "compute", cut_result)
        if measures is None:
            data["unavailable"]["importance"] = reason
        else:
            measures = sorted(list(measures), key=lambda m: -(m.get("fv") or 0.0))
            data["importance"] = {"rows": measures[: opts["topImportance"]], "total": len(measures)}

    if "uncertainty" in wanted:
        result = opts.get("uncertainty")
        reason = None
        if result is None and opts.get("runUncertainty"):
            mc = analysis.get("mc") or {}
            result, reason = try_call(
                "uncertainty", "run", tree, analysis,
                n=min(int(opts.get("uncertaintyN") or MAX_REPORT_MC_N), MAX_REPORT_MC_N),
                seed=mc.get("seed"), time_limit=30,
            )
        if result is None:
            data["unavailable"]["uncertainty"] = reason or "notRun"
        data["uncertainty"] = result

    if "validation" in wanted:
        issues, reason = try_call("lint", "run", tree, analysis, list(session_warnings or []),
                                 mode=mode,
                                 extra={"cutsets": cut_result} if cut_result is not None else None)
        if issues is None:
            data["unavailable"]["validation"] = reason
            issues = list(session_warnings or [])
            for w in getattr(core, "quant_warnings", None) or []:
                issues.append({"severity": "warning", "code": w.get("code"),
                               "nodeId": w.get("nodeId"), "message": "", "params": w.get("params")})
        data["validation"] = [dict(i) for i in issues]

    if "traceability" in wanted:
        rows = []
        for node, _p, _d in iter_nodes(tree):
            trace = node.get("trace") if isinstance(node.get("trace"), dict) else None
            if not trace or not any(v not in (None, "", []) for v in trace.values()):
                continue
            tags = trace.get("tags")
            rows.append({
                "id": node.get("id"), "name": node.get("name", ""),
                "requirementId": trace.get("requirementId") or "",
                "testRef": trace.get("testRef") or "",
                "owner": trace.get("owner") or "",
                "status": trace.get("status") or "",
                "evidence": trace.get("evidence") or "",
                "tags": ", ".join(str(t) for t in tags) if isinstance(tags, list) else (tags or ""),
            })
        data["traceability"] = rows
    data["names"] = names
    return data


# ---- localization -------------------------------------------------------------------------

L10N: Dict[str, Dict[str, str]] = {
    "en": {
        "reportTitle": "Fault Tree Analysis Report",
        "untitled": "Untitled analysis",
        "date": "Analysis date", "mode": "Mode", "generated": "Generated",
        "tool": "Tool", "toolName": "FTA Editor",
        "headline": "Top-event probability",
        "method.mcub": "MCUB (min-cut upper bound)",
        "method.treeWalk": "Tree walk",
        "method": "Method", "alt.treeWalk": "Tree walk value", "alt.mcub": "MCUB value",
        "assumptions": "Assumptions",
        "missionTime": "Mission time: {v} h (display unit: {u})",
        "models": "Quantification models used",
        "model.fixed": "Fixed probability", "model.rate": "Constant failure rate",
        "model.standby": "Periodically tested standby", "model.repairable": "Repairable (steady state)",
        "events": "events", "approximations": "Approximations",
        "approx.none": "No approximations were needed.",
        "approx.PAND_APPROX": "Priority-AND gate {id} approximated as Πp/n! (sequence not modelled).",
        "approx.XOR_ARITY": "XOR gate {id} with a number of inputs other than two is evaluated as odd parity.",
        "approx.STANDBY_LARGE_LT": "Standby event {id}: λτ > 0.2, the λτ/2 approximation is inaccurate.",
        "approx.other": "{code} at {id}.",
        "nonCoherent": "The tree contains XOR gates and is non-coherent; cut-set results treat XOR as OR.",
        "repeated": "Repeated events (shared by several branches): {list}. The headline uses the MCUB.",
        "repeated.none": "No repeated events.",
        "cutLimits": "Cut-set limits: max order {o}, max count {c}, cutoff {x}.",
        "diagram": "Fault tree diagram", "diagram.none": "Diagram not available.",
        "eventTable": "Basic events",
        "cutsets": "Minimal cut sets",
        "cutsets.summary": "Showing {shown} of {total} cut sets.",
        "cutsets.truncated": "The cut-set list was truncated ({by}).",
        "importance": "Importance measures",
        "importance.summary": "Showing {shown} of {total} events, by Fussell-Vesely.",
        "uncertainty": "Uncertainty (Monte Carlo)",
        "uncertainty.notRun": "Monte Carlo was not run for this report.",
        "uncertainty.partial": "Time limit reached: the result is partial.",
        "histogram": "Histogram",
        "validation": "Validation issues", "validation.none": "No issues found.",
        "traceability": "Traceability", "traceability.none": "No traceability data.",
        "unavailable": "This section is not available in this build ({reason}).",
        "etaSkipped": "Event tree (ETA) mode: this section applies to fault trees only.",
        "col.id": "Id", "col.name": "Name", "col.kind": "Kind", "col.model": "Model",
        "col.params": "Parameters", "col.q": "q", "col.calc": "Calculated",
        "col.source": "Source",
        "col.rank": "#", "col.events": "Events", "col.order": "Order",
        "col.prob": "Probability", "col.share": "Share",
        "col.fv": "FV", "col.birnbaum": "Birnbaum", "col.raw": "RAW", "col.rrw": "RRW",
        "col.count": "Cut sets", "col.stat": "Statistic", "col.value": "Value",
        "col.bin": "Bin", "col.from": "From", "col.to": "To", "col.n": "Count",
        "col.severity": "Severity", "col.code": "Code", "col.node": "Node", "col.message": "Message",
        "col.req": "Requirement", "col.test": "Test ref", "col.owner": "Owner",
        "col.status": "Status", "col.evidence": "Evidence", "col.tags": "Tags",
        "stat.pointEstimate": "Point estimate", "stat.mean": "Mean", "stat.median": "Median",
        "stat.p05": "5th percentile", "stat.p95": "95th percentile", "stat.std": "Std. deviation",
        "stat.n": "Samples", "stat.seed": "Seed", "stat.method": "Method",
    },
    "ja": {
        "reportTitle": "故障の木解析レポート",
        "untitled": "無題の解析",
        "date": "解析日", "mode": "モード", "generated": "作成日時",
        "tool": "ツール", "toolName": "FTA Editor",
        "headline": "頂上事象の発生確率",
        "method.mcub": "MCUB(ミニマルカット上限)",
        "method.treeWalk": "ツリー計算",
        "method": "算出方法", "alt.treeWalk": "ツリー計算値", "alt.mcub": "MCUB値",
        "assumptions": "前提条件",
        "missionTime": "ミッション時間: {v} h(表示単位: {u})",
        "models": "使用した定量化モデル",
        "model.fixed": "固定確率", "model.rate": "一定故障率",
        "model.standby": "定期試験の待機系", "model.repairable": "修理系(定常)",
        "events": "事象", "approximations": "近似",
        "approx.none": "近似は使用していません。",
        "approx.PAND_APPROX": "優先AND ゲート {id} は Πp/n! で近似しています(順序は考慮しません)。",
        "approx.XOR_ARITY": "入力が2個以外の XOR ゲート {id} は奇数パリティとして計算しています。",
        "approx.STANDBY_LARGE_LT": "待機系事象 {id}: λτ > 0.2 のため λτ/2 近似の精度が低下します。",
        "approx.other": "{id} の {code}。",
        "nonCoherent": "XOR ゲートを含む非コヒーレントな木です。カットセットでは XOR を OR として扱います。",
        "repeated": "重複事象(複数の枝で共有): {list}。頂上事象は MCUB を使用しています。",
        "repeated.none": "重複事象はありません。",
        "cutLimits": "カットセットの制限: 最大次数 {o}、最大数 {c}、カットオフ {x}。",
        "diagram": "故障の木の図", "diagram.none": "図はありません。",
        "eventTable": "基本事象",
        "cutsets": "ミニマルカットセット",
        "cutsets.summary": "{total} 件中 {shown} 件を表示。",
        "cutsets.truncated": "カットセットの一覧は打ち切られています({by})。",
        "importance": "重要度",
        "importance.summary": "{total} 事象中 {shown} 件を FV 順に表示。",
        "uncertainty": "不確かさ(モンテカルロ)",
        "uncertainty.notRun": "このレポートではモンテカルロ計算を実行していません。",
        "uncertainty.partial": "時間制限に達したため、結果は途中までのものです。",
        "histogram": "ヒストグラム",
        "validation": "検証結果", "validation.none": "問題は見つかりませんでした。",
        "traceability": "トレーサビリティ", "traceability.none": "トレーサビリティ情報はありません。",
        "unavailable": "この版ではこの節は利用できません({reason})。",
        "etaSkipped": "イベントツリー(ETA)モード: この節は故障の木にのみ適用されます。",
        "col.id": "ID", "col.name": "名前", "col.kind": "種類", "col.model": "モデル",
        "col.params": "パラメータ", "col.q": "q", "col.calc": "計算値",
        "col.source": "出典",
        "col.rank": "#", "col.events": "事象", "col.order": "次数",
        "col.prob": "確率", "col.share": "寄与率",
        "col.fv": "FV", "col.birnbaum": "Birnbaum", "col.raw": "RAW", "col.rrw": "RRW",
        "col.count": "カットセット数", "col.stat": "統計量", "col.value": "値",
        "col.bin": "区間", "col.from": "下限", "col.to": "上限", "col.n": "度数",
        "col.severity": "重大度", "col.code": "コード", "col.node": "ノード", "col.message": "メッセージ",
        "col.req": "要求ID", "col.test": "試験参照", "col.owner": "担当者",
        "col.status": "状態", "col.evidence": "エビデンス", "col.tags": "タグ",
        "stat.pointEstimate": "点推定値", "stat.mean": "平均", "stat.median": "中央値",
        "stat.p05": "5パーセンタイル", "stat.p95": "95パーセンタイル", "stat.std": "標準偏差",
        "stat.n": "試行回数", "stat.seed": "シード", "stat.method": "手法",
    },
}


def _tr(lang: str) -> Callable[..., str]:
    table = L10N.get(lang, L10N["en"])

    def t(key: str, **kw: Any) -> str:
        text = table.get(key, L10N["en"].get(key, key))
        for name, value in kw.items():
            text = text.replace("{%s}" % name, str(value))
        return text

    return t


# ---- DOCX building -------------------------------------------------------------------------


def _set_east_asian_font(style, font_name: str) -> None:
    from docx.oxml.ns import qn

    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = rpr.makeelement(qn("w:rFonts"), {})
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), font_name)
    # A theme attribute outranks the explicit name in Word; drop it.
    rfonts.attrib.pop(qn("w:eastAsiaTheme"), None)


def _repeat_header(row) -> None:
    from docx.oxml.ns import qn

    tr_pr = row._tr.get_or_add_trPr()
    el = tr_pr.makeelement(qn("w:tblHeader"), {})
    el.set(qn("w:val"), "true")
    tr_pr.append(el)


def _table(doc, headers: List[str], rows: List[List[Any]]):
    table = doc.add_table(rows=1, cols=len(headers))
    for style_name in ("Light Grid Accent 1", "Table Grid"):
        try:
            table.style = doc.styles[style_name]
            break
        except KeyError:
            continue
    head = table.rows[0]
    for cell, text in zip(head.cells, headers):
        cell.text = str(text)
        for run in cell.paragraphs[0].runs:
            run.bold = True
    _repeat_header(head)
    for row in rows:
        cells = table.add_row().cells
        for cell, value in zip(cells, row):
            cell.text = "" if value is None else str(value)
    return table


def _note(doc, text: str, italic: bool = True):
    para = doc.add_paragraph()
    run = para.add_run(text)
    run.italic = italic
    return para


def _params_text(row: Dict[str, Any], fp: Callable[[Any], str]) -> str:
    parts = []
    for header, label in (("λ (/h)", "λ"), ("T (h)", "T"), ("τ (h)", "τ"), ("μ (/h)", "μ"),
                          ("MTTR (h)", "MTTR")):
        if row.get(header) is not None:
            parts.append("%s=%s" % (label, fp(row[header])))
    return ", ".join(parts)


def _node_label(issue: Dict[str, Any], names: Dict[str, str]) -> str:
    nid = issue.get("nodeId")
    if nid in (None, ""):
        return ""
    name = issue.get("nodeName") or names.get(str(nid))
    return "%s (%s)" % (name, nid) if name else str(nid)


def build_report(doc_data: Dict[str, Any], options: Optional[Dict[str, Any]] = None) -> bytes:
    """The report as ``.docx`` bytes. Needs python-docx."""
    from docx import Document
    from docx.enum.style import WD_STYLE_TYPE
    from docx.shared import Cm, Pt

    options = options or {}
    lang = options.get("lang") or doc_data.get("lang") or "en"
    sf = clamp_sig_figs(options.get("sigFigs", doc_data.get("sigFigs", 3)))
    t = _tr(lang)

    def fp(value: Any) -> str:
        return format_prob(value, sf)

    sections = list(doc_data.get("sections") or DEFAULT_SECTIONS)
    unavailable = doc_data.get("unavailable") or {}
    skipped = set(doc_data.get("skippedEta") or [])
    names = doc_data.get("names") or {}

    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.size = Pt(10)
    if lang == "ja":
        for style in doc.styles:
            try:
                if style.type in (WD_STYLE_TYPE.PARAGRAPH, WD_STYLE_TYPE.CHARACTER,
                                  WD_STYLE_TYPE.TABLE):
                    _set_east_asian_font(style, "Meiryo")
            except Exception:
                continue

    def unavailable_note(section: str) -> bool:
        if section in skipped:
            _note(doc, t("etaSkipped"))
            return True
        if section in unavailable and doc_data.get(section) is None:
            reason = unavailable[section]
            if section == "uncertainty" and reason == "notRun":
                _note(doc, t("uncertainty.notRun"))
            else:
                _note(doc, t("unavailable", reason=reason))
            return True
        return False

    meta = doc_data.get("meta") or {}
    # ---- title page / metadata ----
    doc.add_heading(t("reportTitle"), 0)
    doc.add_heading(meta.get("title") or t("untitled"), 1)
    if "metadata" in sections:
        for key, value in (("date", meta.get("date")), ("mode", meta.get("mode")),
                           ("generated", meta.get("generated")),
                           ("tool", "%s %s" % (t("toolName"), meta.get("version") or TOOL_VERSION))):
            para = doc.add_paragraph()
            para.add_run(t(key) + ": ").bold = True
            para.add_run(str(value or ""))

    # ---- headline ----
    if "headline" in sections:
        doc.add_heading(t("headline"), 1)
        summary = doc_data.get("summary")
        if not unavailable_note("headline") and summary:
            method = summary.get("headlineMethod") or "treeWalk"
            para = doc.add_paragraph()
            run = para.add_run(fp(summary.get("headline")))
            run.bold = True
            run.font.size = Pt(16)
            para.add_run("   [%s]" % t("method." + ("mcub" if method == "mcub" else "treeWalk")))
            alt_key = "treeWalk" if method == "mcub" else "mcub"
            alt = summary.get(alt_key)
            if alt is not None:
                doc.add_paragraph("%s: %s" % (t("alt." + alt_key), fp(alt)))

    # ---- assumptions ----
    if "assumptions" in sections and doc_data.get("assumptions"):
        a = doc_data["assumptions"]
        doc.add_heading(t("assumptions"), 1)
        doc.add_paragraph(t("missionTime", v=a.get("missionTime"), u=a.get("timeUnit")),
                          style="List Bullet")
        cut = a.get("cutsets") or {}
        if cut:
            doc.add_paragraph(t("cutLimits", o=cut.get("maxOrder"), c=cut.get("maxCount"),
                                x=fp(cut.get("cutoff"))), style="List Bullet")
        doc.add_heading(t("models"), 2)
        for m in a.get("models") or []:
            doc.add_paragraph("%s (%s): %s — %d %s" % (
                t("model." + m["model"]), m["model"], m.get("formula") or "", m["count"], t("events")),
                style="List Bullet")
        doc.add_heading(t("approximations"), 2)
        approx = a.get("approximations") or []
        if not approx and not a.get("nonCoherent"):
            doc.add_paragraph(t("approx.none"), style="List Bullet")
        for item in approx:
            code = item.get("code")
            nid = str(item.get("nodeId"))
            label = "%s (%s)" % (names.get(nid, nid), nid)
            key = "approx." + str(code)
            text = t(key, id=label) if key in L10N["en"] else t("approx.other", code=code, id=label)
            doc.add_paragraph(text, style="List Bullet")
        if a.get("nonCoherent"):
            doc.add_paragraph(t("nonCoherent"), style="List Bullet")
        rep = a.get("repeatedEvents") or []
        doc.add_paragraph(t("repeated", list=", ".join(rep)) if rep else t("repeated.none"),
                          style="List Bullet")

    # ---- diagram ----
    if "diagram" in sections:
        doc.add_heading(t("diagram"), 1)
        png = doc_data.get("diagramPng")
        placed = False
        if png:
            try:
                doc.add_picture(io.BytesIO(bytes(png)), width=Cm(16))
                placed = True
            except Exception:
                placed = False
        if not placed:
            _note(doc, t("diagram.none"))

    # ---- events ----
    if "events" in sections and doc_data.get("events") is not None:
        doc.add_heading(t("eventTable"), 1)
        rows = []
        for r in doc_data["events"]:
            rows.append([r.get("Id"), r.get("Name"), r.get("Event kind") or "",
                         r.get("Model") or "", _params_text(r, fp),
                         fp(r.get("Base probability")), fp(r.get("Calculated probability")),
                         r.get("Source") or ""])
        _table(doc, [t("col.id"), t("col.name"), t("col.kind"), t("col.model"), t("col.params"),
                     t("col.q"), t("col.calc"), t("col.source")], rows)

    # ---- cut sets ----
    if "cutsets" in sections:
        doc.add_heading(t("cutsets"), 1)
        cs = doc_data.get("cutsets")
        if not unavailable_note("cutsets") and cs:
            doc.add_paragraph(t("cutsets.summary", shown=cs.get("shown"), total=cs.get("total")))
            if cs.get("truncated"):
                by = cs.get("truncatedBy")
                by = ", ".join(map(str, by)) if isinstance(by, (list, tuple)) else (by or "")
                _note(doc, t("cutsets.truncated", by=by))
            rows = []
            for c in cs.get("rows") or []:
                events = " · ".join(str(e.get("name") or e.get("id")) for e in c.get("events") or [])
                share = c.get("share")
                rows.append([c.get("rank"), events, c.get("order"), fp(c.get("probability")),
                             "" if share is None else "%.1f%%" % (float(share) * 100)])
            _table(doc, [t("col.rank"), t("col.events"), t("col.order"), t("col.prob"),
                         t("col.share")], rows)

    # ---- importance ----
    if "importance" in sections:
        doc.add_heading(t("importance"), 1)
        imp = doc_data.get("importance")
        if not unavailable_note("importance") and imp:
            doc.add_paragraph(t("importance.summary", shown=len(imp.get("rows") or []),
                                total=imp.get("total")))
            rows = [[m.get("id"), m.get("name"), fp(m.get("q")), fp(m.get("fv")),
                     fp(m.get("birnbaum")), fp(m.get("raw")), fp(m.get("rrw")),
                     m.get("cutSetCount")] for m in imp.get("rows") or []]
            _table(doc, [t("col.id"), t("col.name"), t("col.q"), t("col.fv"), t("col.birnbaum"),
                         t("col.raw"), t("col.rrw"), t("col.count")], rows)

    # ---- uncertainty ----
    if "uncertainty" in sections:
        doc.add_heading(t("uncertainty"), 1)
        unc = doc_data.get("uncertainty")
        if not unavailable_note("uncertainty") and unc:
            if unc.get("completed") is False:
                _note(doc, t("uncertainty.partial"))
            rows = []
            for key in ("pointEstimate", "mean", "median", "p05", "p95", "std"):
                if key in unc:
                    rows.append([t("stat." + key), fp(unc.get(key))])
            for key in ("n", "seed", "method"):
                if unc.get(key) is not None:
                    rows.append([t("stat." + key), str(unc.get(key))])
            _table(doc, [t("col.stat"), t("col.value")], rows)
            hist = unc.get("histogram") or {}
            edges, counts = hist.get("edges") or [], hist.get("counts") or []
            if counts and len(edges) == len(counts) + 1:
                doc.add_heading(t("histogram"), 2)
                _table(doc, [t("col.bin"), t("col.from"), t("col.to"), t("col.n")],
                       [[i + 1, fp(edges[i]), fp(edges[i + 1]), counts[i]]
                        for i in range(len(counts))])

    # ---- validation ----
    if "validation" in sections and doc_data.get("validation") is not None:
        doc.add_heading(t("validation"), 1)
        issues = doc_data["validation"]
        if not issues:
            _note(doc, t("validation.none"), italic=False)
        else:
            _table(doc, [t("col.severity"), t("col.code"), t("col.node"), t("col.message")],
                   [[i.get("severity"), i.get("code"), _node_label(i, names), i.get("message") or ""]
                    for i in issues])

    # ---- traceability ----
    if "traceability" in sections and doc_data.get("traceability") is not None:
        doc.add_heading(t("traceability"), 1)
        rows = doc_data["traceability"]
        if not rows:
            _note(doc, t("traceability.none"), italic=False)
        else:
            _table(doc, [t("col.id"), t("col.name"), t("col.req"), t("col.test"), t("col.owner"),
                         t("col.status"), t("col.evidence"), t("col.tags")],
                   [[r["id"], r["name"], r["requirementId"], r["testRef"], r["owner"],
                     r["status"], r["evidence"], r["tags"]] for r in rows])

    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


# ---- server-side diagram fallback ---------------------------------------------------------


def diagram_dot_text(core, sig_figs: int = 3) -> Optional[str]:
    """Compact-style DOT for ``core`` (caller holds the lock), or None."""
    try:
        dd = _module("diagram_dot")
        dot, _id_map = dd.build_dot_text2(core, style="compact", sig_figs=sig_figs)
        return dot
    except Exception:
        return None


def render_png(dot_text: Optional[str]) -> Optional[bytes]:
    """PNG from native Graphviz, or None when it is missing or fails."""
    if not dot_text:
        return None
    try:
        rendering = _module("rendering")
        if not rendering.probe_native_dot():
            return None
        image = rendering.render_native(dot_text, "png")
        return image if image[:8] == PNG_SIGNATURE else None
    except Exception:
        return None
