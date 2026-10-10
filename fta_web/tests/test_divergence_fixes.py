"""
Regression tests for the documented divergences of the vendored core.

Each divergence recorded in fta_web/core/DIVERGENCE.md gets coverage here:

  D1 - AIAgentHandler._get_client() (dead code, always raised AttributeError)
       was deleted.
  D3 - Excel export sibling row layout. Reported as a row-overwrite defect but
       NOT reproducible; no patch was applied. These tests pin the layout so a
       future refactor of write_node() cannot introduce the reported bug.
  D5 - logicGate "NOT" is rejected at validation instead of silently computing
       as OR.
  D7 - the move cycle-check asked the question backwards.
  D27 - a proposal for a node without a gate ("logicGate": "", which is what
       the system prompt's own example shows for an Event) was dropped.
  D28 - a "move" proposal could never be applied: the system prompt did not
       say how to name the new parent.
  D29 - a proposal whose DATA held a link (a nested object) arrived with no
       data, and the system prompt did not say how to write a link change.
"""

import json
import re
import sys
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "core"))

from FTA_Editor_core import FTACore  # noqa: E402
from AI_agent_handler import AIAgentHandler  # noqa: E402


# --------------------------------------------------------------------------
# helpers
# --------------------------------------------------------------------------

def node(node_id, children=None, **kwargs):
    """Build a well-formed FTA node."""
    return {
        "id": node_id,
        "name": kwargs.get("name", node_id),
        "type": kwargs.get("type", "Event"),
        "logicGate": kwargs.get("logicGate", ""),
        "probability": kwargs.get("probability", 0.1),
        "children": children or [],
        "links": [],
        "notes": kwargs.get("notes", ""),
    }


def walk(n):
    """Yield every node in the tree, depth-first."""
    yield n
    for child in n.get("children", []):
        yield from walk(child)


def export_grid(tree, path):
    """Export `tree` to Excel and return {(row, col): first_line_of_cell}."""
    core = FTACore()
    core.set_data(tree)
    ok, err = core.export_to_excel(str(path))
    assert ok, f"export failed: {err}"

    from openpyxl import load_workbook

    ws = load_workbook(str(path)).active
    return {
        (cell.row, cell.column): str(cell.value).split("\n")[0]
        for row in ws.iter_rows()
        for cell in row
        if cell.value
    }


# --------------------------------------------------------------------------
# D1 - dead _get_client() removed
# --------------------------------------------------------------------------

def test_d1_get_client_attribute_is_gone():
    """The dead _get_client method must not exist on the class or instances."""
    assert not hasattr(AIAgentHandler, "_get_client")

    handler = AIAgentHandler()
    assert not hasattr(handler, "_get_client")


def test_d1_live_provider_path_survived():
    """Deleting _get_client must not have touched the live _get_provider path."""
    handler = AIAgentHandler()
    assert hasattr(handler, "_get_provider")
    assert callable(handler._get_provider)
    # _get_provider caches on `provider`, which __init__ *does* initialise --
    # this is exactly what _get_client's `self._client` was missing.
    assert hasattr(handler, "provider")


# --------------------------------------------------------------------------
# D3 - Excel export sibling row layout
# --------------------------------------------------------------------------

def test_d3_deep_sibling_tree_loses_no_nodes(tmp_path):
    """A tree whose first sibling branch has grandchildren must export intact."""
    tree = node(
        "ROOT",
        [
            node(
                "S1",
                [
                    node("S1a", [node("S1a1"), node("S1a2")]),
                    node("S1b", [node("S1b1")]),
                ],
            ),
            node("S2", [node("S2a")]),
            node("S3"),
        ],
        type="Root",
    )

    grid = export_grid(tree, tmp_path / "deep.xlsx")

    expected_names = {n["name"] for n in walk(tree)}
    written_names = set(grid.values())

    missing = expected_names - written_names
    assert not missing, f"nodes lost from the sheet: {sorted(missing)}"

    # No node may be dropped by a same-cell overwrite either.
    assert len(grid) == len(list(walk(tree)))


def test_d3_deep_sibling_tree_exact_layout(tmp_path):
    """Pin the hierarchical layout: depth -> column, subtrees never collide.

    S2 must start on row 4, *below* S1's three-row subtree -- this is the exact
    scenario the reported defect claimed would overwrite rows.
    """
    tree = node(
        "ROOT",
        [
            node(
                "S1",
                [
                    node("S1a", [node("S1a1"), node("S1a2")]),
                    node("S1b", [node("S1b1")]),
                ],
            ),
            node("S2", [node("S2a")]),
        ],
        type="Root",
    )

    grid = export_grid(tree, tmp_path / "layout.xlsx")

    assert grid == {
        (1, 1): "ROOT",
        (1, 2): "S1",
        (1, 3): "S1a",
        (1, 4): "S1a1",
        (2, 4): "S1a2",
        (3, 3): "S1b",
        (3, 4): "S1b1",
        (4, 2): "S2",
        (4, 3): "S2a",
    }


def test_d3_shallow_tree_output_unchanged(tmp_path):
    """A tree with no multi-level sibling descendants keeps its original layout."""
    tree = node("ROOT", [node("X"), node("Y"), node("Z")], type="Root")

    grid = export_grid(tree, tmp_path / "shallow.xlsx")

    assert grid == {
        (1, 1): "ROOT",
        (1, 2): "X",
        (2, 2): "Y",
        (3, 2): "Z",
    }


def test_d3_cell_content_and_styling_preserved(tmp_path):
    """Content format, per-depth fills, bold root, wrap/top align, widths, heights."""
    tree = node(
        "ROOT",
        [node("child", probability=0.25, type="Gate", logicGate="AND", notes="a note")],
        type="Root",
    )
    core = FTACore()
    core.set_data(tree)
    core.recalculate_probabilities()

    path = tmp_path / "style.xlsx"
    ok, err = core.export_to_excel(str(path))
    assert ok, err

    from openpyxl import load_workbook

    ws = load_workbook(str(path)).active

    root_cell = ws.cell(row=1, column=1)
    child_cell = ws.cell(row=1, column=2)

    # Content format: name, then "(Type: ..., P: ..., Calc: ..., Gate: ...)",
    # then "Notes: ...".
    lines = child_cell.value.split("\n")
    assert lines[0] == "child"
    assert lines[1].startswith("(Type: Gate") and lines[1].endswith(")")
    assert "P: 0.25" in lines[1]
    assert "Gate: AND" in lines[1]
    assert lines[2] == "Notes: a note"

    # Bold root font, plain child.
    assert root_cell.font.bold is True
    assert not child_cell.font.bold

    # wrap_text / vertical top on both.
    for cell in (root_cell, child_cell):
        assert cell.alignment.wrap_text is True
        assert cell.alignment.vertical == "top"

    # Per-depth fill colours.
    assert root_cell.fill.start_color.rgb.endswith("E6F3FF")
    assert child_cell.fill.start_color.rgb.endswith("FFF4E6")

    # Auto column widths capped at 50, row heights 45.
    for dim in ws.column_dimensions.values():
        assert dim.width <= 50
    assert ws.row_dimensions[1].height == 45


def test_d3_links_rendered_in_cell(tmp_path):
    """The 'Links: ...' line of the cell format is preserved."""
    target = node("tgt", name="Target Node")
    source = node("src", name="Source Node")
    source["links"] = [{"target_id": "tgt", "relation": "AND"}]
    tree = node("ROOT", [source, target], type="Root")

    core = FTACore()
    core.set_data(tree)
    path = tmp_path / "links.xlsx"
    ok, err = core.export_to_excel(str(path))
    assert ok, err

    from openpyxl import load_workbook

    ws = load_workbook(str(path)).active
    assert "Links: AND→Target Node" in ws.cell(row=1, column=2).value


# --------------------------------------------------------------------------
# D5 - NOT gates are rejected
# --------------------------------------------------------------------------

def _tree_with_gate(gate):
    return node(
        "root",
        [node("child_a", probability=0.2), node("child_b", probability=0.3)],
        type="Root",
        logicGate=gate,
        probability=1.0,
    )


def test_d5_not_gate_rejected_by_verify_updated_fta_json():
    handler = AIAgentHandler()
    ok, err = handler.verify_updated_fta_json(_tree_with_gate("NOT"))

    assert ok is False
    assert err is not None
    # The message must name the offending node.
    assert "root" in err
    assert "NOT" in err


def test_d5_not_gate_rejected_on_a_nested_node_names_that_node():
    handler = AIAgentHandler()
    tree = node(
        "root",
        [node("branch", [node("leaf", probability=0.4)], logicGate="NOT", type="Gate")],
        type="Root",
        probability=1.0,
    )

    ok, err = handler.verify_updated_fta_json(tree)

    assert ok is False
    assert "branch" in err, f"error must name the offending node, got: {err!r}"


@pytest.mark.parametrize("gate", ["AND", "OR", ""])
def test_d5_valid_gates_still_accepted(gate):
    handler = AIAgentHandler()
    ok, err = handler.verify_updated_fta_json(_tree_with_gate(gate))

    assert ok is True, f"gate {gate!r} should validate, got error: {err!r}"
    assert err is None


def test_d5_not_gate_rejected_by_validate_node_data():
    handler = AIAgentHandler()
    ok, err = handler._validate_node_data(
        {"id": "gate_3", "name": "Some Gate", "type": "Gate", "logicGate": "NOT"}
    )

    assert ok is False
    assert "gate_3" in err
    assert "NOT" in err


@pytest.mark.parametrize("gate", ["AND", "OR"])
def test_d5_validate_node_data_still_accepts_and_or(gate):
    handler = AIAgentHandler()
    ok, err = handler._validate_node_data(
        {"id": "gate_3", "name": "Some Gate", "type": "Gate", "logicGate": gate}
    )

    assert ok is True, f"gate {gate!r} should validate, got: {err!r}"


def test_d5_prompts_never_advertise_not_as_valid():
    """The AI must not be told NOT is an accepted logicGate value."""
    assert "AND|OR|NOT" not in AIAgentHandler.FTA_JSON_SCHEMA
    assert "AND|OR" in AIAgentHandler.FTA_JSON_SCHEMA
    assert "AND|OR|NOT" not in AIAgentHandler.SYSTEM_PROMPT


def test_d5_probability_engine_unchanged_for_and_or():
    """D5 was fixed at validation only; the engine must be untouched."""
    core = FTACore()
    core.set_data(_tree_with_gate("AND"))
    core.recalculate_probabilities()
    assert core.get_data()["calculatedProbability"] == pytest.approx(0.2 * 0.3)

    core = FTACore()
    core.set_data(_tree_with_gate("OR"))
    core.recalculate_probabilities()
    assert core.get_data()["calculatedProbability"] == pytest.approx(
        1 - (1 - 0.2) * (1 - 0.3)
    )


# ---------------------------------------------------------------------------
# D7 -- move cycle-check asked the question backwards
#
# At baseline, _would_create_circular_reference delegated to
# _is_descendant_of(core, node_id, potential_parent_id) -- "is the node already
# inside the target's subtree" -- which is the inverse of the guard a move needs.
# It rejected every legal move and permitted every corrupting one.
# ---------------------------------------------------------------------------


def _lineage():
    """root > parent > child > grandchild, for exercising move legality."""

    def n(nid, kids=()):
        return {
            "id": nid, "name": nid, "type": "Event", "probability": 0.5,
            "logicGate": "OR", "notes": "", "links": [], "children": list(kids),
        }

    core = FTACore()
    core.set_data(n("root", [n("parent", [n("child", [n("grandchild")])])]))
    return core


@pytest.mark.parametrize(
    "node_id,new_parent_id",
    [
        ("grandchild", "root"),    # legal promotion, two levels up
        ("child", "parent"),       # reorder under the current parent
        ("grandchild", "parent"),  # legal promotion, one level up
    ],
)
def test_d7_legal_moves_are_allowed(node_id, new_parent_id):
    handler = AIAgentHandler()
    assert handler._would_create_circular_reference(
        _lineage(), node_id, new_parent_id
    ) is False


@pytest.mark.parametrize(
    "node_id,new_parent_id",
    [
        ("parent", "child"),        # into its own child
        ("parent", "grandchild"),   # into its own grandchild
        ("root", "parent"),         # root into its own subtree
        ("child", "child"),         # into itself
    ],
)
def test_d7_cycle_creating_moves_are_rejected(node_id, new_parent_id):
    handler = AIAgentHandler()
    assert handler._would_create_circular_reference(
        _lineage(), node_id, new_parent_id
    ) is True


def test_d7_returns_a_real_bool_not_a_node_dict():
    """Baseline returned `find(...) and is_descendant(...)`, leaking a node dict."""
    result = AIAgentHandler()._would_create_circular_reference(
        _lineage(), "grandchild", "root"
    )
    assert isinstance(result, bool)


def test_d7_unknown_node_does_not_raise():
    assert AIAgentHandler()._would_create_circular_reference(
        _lineage(), "no_such_node", "root"
    ) is False


# ---------------------------------------------------------------------------
# D27 -- a proposal for a node without a gate was dropped
#
# SYSTEM_PROMPT shows the model how to add an Event: with "logicGate": "".
# _parse_proposed_changes accepted only "AND" and "OR", so a model that
# followed the example had every such proposal discarded without a word
# (seen with Claude: 12 SUGGESTION blocks in the reply, 8 proposals offered).
# ---------------------------------------------------------------------------


def _prompt_example():
    """The two SUGGESTION blocks SYSTEM_PROMPT itself gives the model."""
    match = re.search(
        r"EXAMPLE of adding nested nodes[^\n]*\n(SUGGESTION:.*?)\n\s*(?:EXAMPLE|CRITICAL RULES)",
        AIAgentHandler.SYSTEM_PROMPT,
        re.DOTALL,
    )
    assert match, "SYSTEM_PROMPT no longer carries its worked example"
    return match.group(1)


def _proposal(action="add", target="root", **data):
    full = {"id": "corrosion", "name": "Corrosion", "type": "Event", "probability": 0.05}
    full.update(data)
    return (
        "SUGGESTION: Corrosion\nDESCRIPTION: missing root cause.\n"
        "ACTION: %s\nTARGET: %s\nDATA: %s\n" % (action, target, json.dumps(full))
    )


def test_d27_the_system_prompts_own_example_parses_in_full():
    changes = AIAgentHandler()._parse_proposed_changes(_prompt_example())

    assert [(c.change_type, c.target_id, c.data["id"]) for c in changes] == [
        ("add", "root", "material_deg"),
        ("add", "material_deg", "corrosion"),
    ]
    assert changes[1].data["logicGate"] == ""


def test_d27_the_system_prompts_own_example_applies_as_a_nested_pair():
    handler = AIAgentHandler()
    core = FTACore()
    core.set_data(node("root", type="Root", logicGate="OR"))

    for change in handler._parse_proposed_changes(_prompt_example()):
        assert handler.validate_change(change) == (True, None)
        ok, message = handler.apply_change_to_fta(core, change)
        assert ok, message

    (gate,) = core.get_data()["children"]
    assert (gate["id"], gate["type"], gate["logicGate"]) == ("material_deg", "Gate", "OR")
    (event,) = gate["children"]
    assert (event["id"], event["type"], event["logicGate"]) == ("corrosion", "Event", "")


@pytest.mark.parametrize("action", ["add", "edit"])
@pytest.mark.parametrize("gate", ["", None])
def test_d27_a_proposal_without_a_gate_is_kept(action, gate):
    changes = AIAgentHandler()._parse_proposed_changes(_proposal(action, logicGate=gate))

    assert len(changes) == 1
    # null means the same as "": no gate. It is stored as text either way.
    assert changes[0].data["logicGate"] == ""


@pytest.mark.parametrize("gate", ["NOT", "XOR", "and", 0, 1, False, True, []])
def test_d27_an_unsupported_gate_is_still_dropped(gate):
    """D5 stands: only AND and OR compute, so nothing else may be proposed."""
    assert AIAgentHandler()._parse_proposed_changes(_proposal(logicGate=gate)) == []


def test_d27_validate_node_data_accepts_a_node_without_a_gate():
    ok, err = AIAgentHandler()._validate_node_data(
        {"id": "corrosion", "name": "Corrosion", "type": "Event", "logicGate": ""}
    )

    assert (ok, err) == (True, None)


@pytest.mark.parametrize("gate", ["XOR", "and", None, 0, False, [], {}])
def test_d27_validate_node_data_still_rejects_an_unsupported_gate(gate):
    ok, err = AIAgentHandler()._validate_node_data(
        {"id": "gate_3", "name": "Some Gate", "type": "Gate", "logicGate": gate}
    )

    assert ok is False
    assert "gate_3" in err


@pytest.mark.parametrize("gate", ["", None])
def test_d27_an_edit_without_a_gate_leaves_the_nodes_gate_alone(gate):
    """Writing "" over an AND would turn it into an OR without saying so
    (the engine computes every gate that is not AND as OR)."""
    handler = AIAgentHandler()
    core = FTACore()
    core.set_data(node("root", type="Root", logicGate="OR", children=[
        node("both", type="Gate", logicGate="AND", children=[
            node("a", probability=0.2), node("b", probability=0.3)]),
    ]))
    core.recalculate_probabilities()
    (change,) = handler._parse_proposed_changes(_proposal(
        "edit", target="both", id="both", name="Both fail", type="Gate",
        probability=0.5, logicGate=gate, notes="renamed"))

    ok, message = handler.apply_change_to_fta(core, change)
    core.recalculate_probabilities()

    assert ok, message
    both = core.find_node_by_id("both")
    assert (both["name"], both["notes"], both["logicGate"]) == ("Both fail", "renamed", "AND")
    assert "logicGate" not in message
    assert both["calculatedProbability"] == pytest.approx(0.2 * 0.3)


def test_d27_an_edit_that_only_carries_an_empty_gate_changes_nothing():
    handler = AIAgentHandler()
    core = FTACore()
    core.set_data(node("root", type="Root", logicGate="AND"))
    (change,) = handler._parse_proposed_changes(
        _proposal("edit", target="root", id="root", name="root", logicGate=""))
    change.data = {"id": "root", "logicGate": ""}

    ok, message = handler.apply_change_to_fta(core, change)

    assert (ok, message) == (False, "No valid fields to update")
    assert core.get_data()["logicGate"] == "AND"


# ---------------------------------------------------------------------------
# D28 -- a "move" proposal could never be applied
#
# SYSTEM_PROMPT offered ACTION: move without saying how to name the new
# parent, and apply_change_to_fta read only data["parent_id"]. Claude wrote
# "new_parent_id" every time (6 of 6 moves in two replies), so every move was
# refused with "Move requires 'parent_id' in data" -- and a restructuring
# batch was left half applied: the new gate was added and stayed empty.
# ---------------------------------------------------------------------------

#: One block of a real reply (Claude, 2026-10-10), verbatim.
MODEL_MOVE = (
    "SUGGESTION: 常用ポンプ停止を中間ゲートの下へ移動\n"
    "DESCRIPTION: 「常用ポンプが止まる」を independent_both_fail の子に移します。\n"
    "ACTION: move\n"
    "TARGET: root_0\n"
    'DATA: {"id": "root_0", "name": "常用ポンプが止まる", "type": "Gate", "probability": 1.0, '
    '"logicGate": "OR", "notes": "移動先: independent_both_fail", '
    '"new_parent_id": "independent_both_fail"}\n'
)


def _prompt_examples():
    """Every SUGGESTION block SYSTEM_PROMPT gives the model as an example."""
    match = re.search(
        r"EXAMPLE of adding nested nodes[^\n]*\n(SUGGESTION:.*?)\n\s*CRITICAL RULES",
        AIAgentHandler.SYSTEM_PROMPT,
        re.DOTALL,
    )
    assert match, "SYSTEM_PROMPT no longer carries its worked examples"
    return match.group(1)


def _two_branches():
    core = FTACore()
    core.set_data(node("root", type="Root", logicGate="OR", children=[
        node("root_0", type="Gate", logicGate="OR", children=[node("root_0_0")]),
        node("independent_both_fail", type="Gate", logicGate="AND"),
    ]))
    return core


def _move(target, data):
    return (
        "SUGGESTION: m\nDESCRIPTION: d\nACTION: move\nTARGET: %s\nDATA: %s\n" % (target, data)
    )


def test_d28_the_system_prompt_shows_a_move_with_its_new_parent():
    blocks = AIAgentHandler()._parse_proposed_changes(_prompt_examples())

    moves = [c for c in blocks if c.change_type == "move"]
    assert len(moves) == 1, "SYSTEM_PROMPT must show one worked move"
    assert moves[0].data == {"parent_id": "material_deg"}
    format_section = AIAgentHandler.SYSTEM_PROMPT.split("EXAMPLE of adding")[0]
    assert re.search(r'^- move: [^\n]*DATA is \{"parent_id": "\w+"\}', format_section, re.M), (
        "the format section, not only the example, must name the key"
    )


def test_d28_the_system_prompts_examples_apply_in_order():
    """What the model is shown must be something the handler can do."""
    handler = AIAgentHandler()
    core = FTACore()
    core.set_data(node("root", type="Root", logicGate="OR", children=[node("fatigue")]))

    changes = handler._parse_proposed_changes(_prompt_examples())
    for change in changes:
        assert handler.validate_change(change) == (True, None)
        ok, message = handler.apply_change_to_fta(core, change)
        assert ok, message

    assert [(c.change_type, c.target_id) for c in changes] == [
        ("add", "root"), ("add", "material_deg"), ("move", "fatigue"), ("edit", "fatigue"),
    ]
    (gate,) = core.get_data()["children"]
    assert gate["id"] == "material_deg"
    assert [child["id"] for child in gate["children"]] == ["corrosion", "fatigue"]


def test_d28_a_move_as_the_model_wrote_it_is_applied():
    handler = AIAgentHandler()
    core = _two_branches()
    (change,) = handler._parse_proposed_changes(MODEL_MOVE)

    ok, message = handler.apply_change_to_fta(core, change)

    assert ok, message
    root = core.get_data()
    assert [child["id"] for child in root["children"]] == ["independent_both_fail"]
    (moved,) = root["children"][0]["children"]
    assert moved["id"] == "root_0"
    assert [child["id"] for child in moved["children"]] == ["root_0_0"], "the subtree moves with it"
    # A move changes where the node hangs and nothing else.
    assert (moved["name"], moved["logicGate"], moved["notes"]) == ("root_0", "OR", "")


def test_d28_parent_id_wins_when_both_keys_are_given():
    handler = AIAgentHandler()
    core = _two_branches()
    (change,) = handler._parse_proposed_changes(_move(
        "root_0_0", '{"parent_id": "independent_both_fail", "new_parent_id": "root"}'))

    ok, message = handler.apply_change_to_fta(core, change)

    assert ok, message
    target = core.find_node_by_id("independent_both_fail")
    assert [child["id"] for child in target["children"]] == ["root_0_0"]


@pytest.mark.parametrize("data", ['{}', '{"id": "root_0"}', '{"parent_id": ""}',
                                  '{"new_parent_id": null}', '{"parent": "root"}'])
def test_d28_a_move_that_names_no_parent_is_refused_and_says_which_key(data):
    handler = AIAgentHandler()
    core = _two_branches()
    before = json.dumps(core.get_data(), sort_keys=True)
    (change,) = handler._parse_proposed_changes(_move("root_0", data))

    ok, message = handler.apply_change_to_fta(core, change)

    assert (ok, message) == (False, "Move requires 'parent_id' in data")
    assert json.dumps(core.get_data(), sort_keys=True) == before


def test_d28_a_move_into_its_own_subtree_is_still_refused():
    """D7 stands whichever key names the parent."""
    handler = AIAgentHandler()
    core = _two_branches()
    (change,) = handler._parse_proposed_changes(_move("root_0", '{"new_parent_id": "root_0_0"}'))

    ok, message = handler.apply_change_to_fta(core, change)

    assert (ok, message) == (False, "Cannot move: would create circular reference")


# --------------------------------------------------------------------------
# D29 -- a proposal whose DATA held a link arrived with no data
# --------------------------------------------------------------------------

LINK_EDIT = (
    "SUGGESTION: Link the secondary seal to the cold\n"
    "DESCRIPTION: The cold acts on both O-rings\n"
    "ACTION: edit\n"
    "TARGET: n_2\n"
    'DATA: {"id": "n_2", "name": "Secondary seal leak", '
    '"links": [{"target_id": "n_1", "relation": "OR"}]}\n'
)

LINK_ADD = (
    "SUGGESTION: Add ice\n"
    "DESCRIPTION: Ice on the pad\n"
    "ACTION: add\n"
    "TARGET: root\n"
    'DATA: {"id": "ice", "name": "Ice", "type": "Event", "probability": 0.1, "logicGate": "", '
    '"notes": "", "links": [{"target_id": "n_1", "relation": "AND"}, '
    '{"target_id": "n_2", "relation": "OR"}]}\n'
)


def test_d29_a_suggestion_with_a_link_keeps_its_data():
    (change,) = AIAgentHandler()._parse_proposed_changes(LINK_EDIT)

    assert (change.change_type, change.target_id) == ("edit", "n_2")
    assert change.data == {
        "id": "n_2", "name": "Secondary seal leak",
        "links": [{"target_id": "n_1", "relation": "OR"}],
    }


def test_d29_every_link_of_an_added_node_is_read():
    (change,) = AIAgentHandler()._parse_proposed_changes(LINK_ADD)

    assert change.data["links"] == [
        {"target_id": "n_1", "relation": "AND"}, {"target_id": "n_2", "relation": "OR"},
    ]
    assert change.data["probability"] == 0.1 and change.data["logicGate"] == ""


def test_d29_blocks_after_a_link_are_still_read_one_by_one():
    """The whole object is read from its own brace, and no further."""
    plain = (
        "SUGGESTION: Rename\nDESCRIPTION: d\nACTION: edit\nTARGET: n_1\n"
        'DATA: {"id": "n_1", "name": "Cold weather"}\n'
    )
    changes = AIAgentHandler()._parse_proposed_changes(LINK_EDIT + "\n" + plain + "\n" + LINK_ADD)

    assert [c.target_id for c in changes] == ["n_2", "n_1", "root"]
    assert changes[1].data == {"id": "n_1", "name": "Cold weather"}
    assert len(changes[2].data["links"]) == 2


def test_d29_a_nested_object_is_held_to_the_same_checks():
    """Read whole, the DATA goes through the parser's field checks: a NOT
    gate is still dropped (D5), nested object or not."""
    bad = LINK_ADD.replace('"logicGate": ""', '"logicGate": "NOT"')

    assert AIAgentHandler()._parse_proposed_changes(bad) == []


def test_d29_data_that_is_not_json_is_still_an_empty_suggestion():
    broken = LINK_EDIT.replace('"relation": "OR"}]}', '"relation": "OR"}')
    (change,) = AIAgentHandler()._parse_proposed_changes(broken)

    assert change.data == {}


def test_d29_the_system_prompt_says_how_to_change_links():
    prompt = AIAgentHandler.SYSTEM_PROMPT
    format_section = prompt.split("EXAMPLE of adding")[0]

    assert '"links": [{"target_id": "id_of_an_existing_node", "relation": "AND|OR"}]' in format_section
    assert "COMPLETE new list" in format_section and "[] removes them all" in format_section
    links = [c for c in AIAgentHandler()._parse_proposed_changes(_prompt_examples())
             if "links" in c.data]
    assert len(links) == 1, "SYSTEM_PROMPT must show one worked link change"
    assert (links[0].change_type, links[0].target_id) == ("edit", "fatigue")
    assert links[0].data["links"] == [{"target_id": "corrosion", "relation": "OR"}]
