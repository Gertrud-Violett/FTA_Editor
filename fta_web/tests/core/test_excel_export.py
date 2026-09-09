#!/usr/bin/env python3
"""
Test script for hierarchical Excel export
Creates a test Excel file to verify the new nested column structure
"""
import sys
from pathlib import Path

# Add the vendored core directory to path (fta_web/core).
# tests/core/<file>.py -> parents[0]=core, [1]=tests, [2]=fta_web
sys.path.insert(0, str(Path(__file__).resolve().parents[2] / "core"))

from FTA_Editor_core import FTACore


def test_hierarchical_export():
    """Test the hierarchical Excel export"""
    import tempfile
    import os

    print("Testing Hierarchical Excel Export")
    print("=" * 70)

    # Create core instance and load sample data
    core = FTACore()

    # Find sampleFTA.json in the vendored fta_web/examples directory
    sample_data_path = Path(__file__).resolve().parents[2] / "examples" / "sampleFTA.json"
    assert sample_data_path.exists(), (
        f"Vendored sample data not found at {sample_data_path}"
    )

    success, error = core.load_from_json(str(sample_data_path))
    assert success, f"Failed to load sample data from {sample_data_path}: {error}"

    print(f"✓ Loaded sampleFTA.json from {sample_data_path}")

    # Recalculate probabilities
    core.recalculate_probabilities()
    print("✓ Calculated probabilities")

    # Export to temp file
    with tempfile.NamedTemporaryFile(delete=False, suffix='.xlsx') as f:
        output_file = f.name

    try:
        success, error = core.export_to_excel(output_file)
        assert success, f"core.export_to_excel() failed: {error}"

        print(f"✓ Exported to {output_file}")

        # Verify the file was created
        assert Path(output_file).exists(), (
            f"export_to_excel() reported success but no file exists at {output_file}"
        )

        # Read back and verify structure.
        #
        # Deliberately NOT wrapped in a broad try/except: AssertionError is a
        # subclass of Exception, so catching it here would silently swallow the
        # very failures these assertions exist to surface. This is precisely how
        # defect 6 hid failures behind `return False`.
        from openpyxl import load_workbook

        wb = load_workbook(output_file)
        ws = wb.active

        print("\n" + "=" * 70)
        print("Excel Structure Preview:")
        print("=" * 70)

        # Show first 15 rows with column positions
        for row_idx in range(1, min(16, ws.max_row + 1)):
            row_data = []
            for col_idx in range(1, min(6, ws.max_column + 1)):  # Show first 5 columns
                cell = ws.cell(row=row_idx, column=col_idx)
                if cell.value:
                    # Truncate long values
                    val = str(cell.value).split('\n')[0]  # Just show first line
                    if len(val) > 25:
                        val = val[:22] + "..."
                    col_letter = chr(64 + col_idx)  # A, B, C, etc.
                    row_data.append(f"{col_letter}{row_idx}: {val}")

            if row_data:
                print(f"Row {row_idx:2d}: " + " | ".join(row_data))

        print("\n" + "=" * 70)
        print(f"Total rows: {ws.max_row}")
        print(f"Total columns: {ws.max_column}")
        print("=" * 70)

        # Verify hierarchical structure
        print("\nVerifying hierarchical structure:")

        # Root should be in column A (column 1)
        root_cell = ws.cell(row=1, column=1)
        assert root_cell.value and "Root Event" in str(root_cell.value), (
            "Root event not in expected position A1. "
            f"Found: {root_cell.value!r}"
        )
        print("✓ Root event found in column A, row 1")

        # Check for children in column B
        first_child_row = None
        for row_idx in range(1, ws.max_row + 1):
            if ws.cell(row=row_idx, column=2).value:
                first_child_row = row_idx
                break

        assert first_child_row is not None, (
            "No children found in column B -- the hierarchical export did not "
            f"nest any child nodes one level in (sheet is {ws.max_row} rows x "
            f"{ws.max_column} cols)"
        )
        child_val = str(ws.cell(row=first_child_row, column=2).value).split(chr(10))[0][:30]
        print(f"✓ Found child in column B, row {first_child_row}: {child_val}")

        # Check for grandchildren in column C
        first_grandchild_row = None
        for row_idx in range(1, ws.max_row + 1):
            if ws.cell(row=row_idx, column=3).value:
                first_grandchild_row = row_idx
                break

        assert first_grandchild_row is not None, (
            "No grandchildren found in column C -- the hierarchical export did "
            f"not nest any second-level nodes (sheet is {ws.max_row} rows x "
            f"{ws.max_column} cols)"
        )
        gc_val = str(ws.cell(row=first_grandchild_row, column=3).value).split(chr(10))[0][:30]
        print(f"✓ Found grandchild in column C, row {first_grandchild_row}: {gc_val}")

        print("\n" + "=" * 70)
        print("✅ ALL TESTS PASSED")
        print("=" * 70)
    finally:
        if os.path.exists(output_file):
            os.unlink(output_file)


if __name__ == "__main__":
    # An assertion failure propagates and exits non-zero via the traceback.
    test_hierarchical_export()
    sys.exit(0)
