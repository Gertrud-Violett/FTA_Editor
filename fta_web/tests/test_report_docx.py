"""The DOCX report: data collection, the document, and POST /api/report/docx."""
import base64
import io
import json
import struct
import zlib

import pytest

from fta_web import report_docx
from fta_web.state import get_state, reset_state

import sys  # noqa: E402
from pathlib import Path  # noqa: E402

if str(Path(__file__).resolve().parent) not in sys.path:
    sys.path.insert(0, str(Path(__file__).resolve().parent))
from test_excel_events import make_core, sample_tree  # noqa: E402


def tiny_png(width=4, height=3):
    """A valid RGB PNG built by hand (no PIL)."""
    def chunk(kind, data):
        body = kind + data
        return struct.pack(">I", len(data)) + body + struct.pack(">I", zlib.crc32(body) & 0xFFFFFFFF)

    raw = b"".join(b"\x00" + b"\xff\x00\x00" * width for _ in range(height))
    return (report_docx.PNG_SIGNATURE
            + chunk(b"IHDR", struct.pack(">IIBBBBB", width, height, 8, 2, 0, 0, 0))
            + chunk(b"IDAT", zlib.compress(raw))
            + chunk(b"IEND", b""))


# ---- fakes for the other workstreams' functions ------------------------------------


def fake_cutsets(tree, analysis, limits=None):
    return {
        "cutSets": [
            {"rank": 1, "events": [{"id": "c", "name": "Sensor", "q": 2e-7}], "order": 1,
             "probability": 0.0123, "share": 0.6},
            {"rank": 2, "events": [{"id": "a", "name": "Pump λ", "q": 1e-2},
                                   {"id": "b", "name": "Valve", "q": 1e-2}],
             "order": 2, "probability": 1e-4, "share": 0.4},
        ],
        "total": 2, "truncated": True, "truncatedBy": ["maxOrder"],
        "mcub": 0.0124, "rareEvent": 0.0125, "treeWalk": 0.0124, "repeatedEvents": [],
        "nonCoherent": False, "approximations": [], "warnings": [], "elapsedMs": 1,
    }


def fake_importance(result):
    return [
        {"id": "a", "name": "Pump λ", "q": 1e-2, "fv": 0.1, "birnbaum": 0.01, "raw": 2.0,
         "rrw": 1.1, "cutSetCount": 1},
        {"id": "c", "name": "Sensor", "q": 2e-7, "fv": 0.9, "birnbaum": 0.9, "raw": 9.0,
         "rrw": 10.0, "cutSetCount": 1},
    ]


def fake_uncertainty(tree, analysis, n=None, seed=None, time_limit=30, bins=40):
    """The shape of uncertainty.run: ``completed`` is a sample count and a
    time-capped run says so in ``truncatedByTime``."""
    fake_uncertainty.calls.append(n)
    fake_uncertainty.time_limits.append(time_limit)
    completed = fake_uncertainty.completed if fake_uncertainty.completed is not None else n
    return {"mean": 1.5e-3, "median": 1.2e-3, "p05": 4e-4, "p95": 3.3e-3, "std": 1e-3,
            "pointEstimate": 1.3e-3, "histogram": {"edges": [0, 1e-3, 2e-3], "counts": [7, 3]},
            "requested": n, "completed": completed, "truncatedByTime": completed != n,
            "seed": 42, "method": "tree"}


fake_uncertainty.calls = []
fake_uncertainty.time_limits = []
fake_uncertainty.completed = None


def fake_lint(tree, analysis, session_warnings=(), mode="FTA", extra=None):
    fake_lint.calls.append({"mode": mode, "extra": extra})
    return [{"severity": "error", "code": "DANGLING_LINK", "nodeId": "b",
             "message": "Link target missing", "params": {}}] + list(session_warnings)


def fake_summary(tree, analysis=None):
    return {"treeWalk": 0.0124, "mcub": 0.0131, "rareEvent": 0.0135, "headline": 0.0131,
            "headlineMethod": "mcub", "repeatedEvents": ["b"], "nonCoherent": True,
            "approximations": [{"code": "PAND_APPROX", "nodeId": "g1", "params": {}}],
            "truncated": False}


fake_lint.calls = []


@pytest.fixture
def fakes(monkeypatch):
    import fta_web.cutsets
    import fta_web.engine
    import fta_web.importance
    import fta_web.lint
    import fta_web.uncertainty

    monkeypatch.setattr(fta_web.cutsets, "compute", fake_cutsets, raising=False)
    monkeypatch.setattr(fta_web.importance, "compute", fake_importance, raising=False)
    monkeypatch.setattr(fta_web.uncertainty, "run", fake_uncertainty, raising=False)
    monkeypatch.setattr(fta_web.lint, "run", fake_lint, raising=False)
    monkeypatch.setattr(fta_web.engine, "summary", fake_summary)
    fake_uncertainty.calls = []
    fake_uncertainty.time_limits = []
    fake_uncertainty.completed = None


ALL = list(report_docx.SECTIONS)


def text_of(doc):
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(c.text for c in row.cells))
    return "\n".join(parts)


# ---- adapters / data ----------------------------------------------------------------


def test_try_call_reports_missing_functions():
    result, reason = report_docx.try_call("cutsets", "definitely_not_there")
    assert result is None and reason == "unavailable"
    result, reason = report_docx.try_call("no_such_module_xyz", "compute")
    assert result is None and reason == "unavailable"


def test_try_call_turns_exceptions_into_reasons(monkeypatch):
    import fta_web.cutsets

    def boom(*a, **k):
        raise NotImplementedError

    monkeypatch.setattr(fta_web.cutsets, "compute", boom, raising=False)
    assert report_docx.try_call("cutsets", "compute", {}, {}) == (None, "unavailable")

    def fail(*a, **k):
        raise RuntimeError("bad tree")

    monkeypatch.setattr(fta_web.cutsets, "compute", fail, raising=False)
    result, reason = report_docx.try_call("cutsets", "compute", {}, {})
    assert result is None and reason.startswith("failed") and "bad tree" in reason


def test_collect_with_fakes(fakes):
    core = make_core()
    before = json.dumps(core.get_data(), sort_keys=True)
    data = report_docx.collect_report_data(
        core, [{"severity": "warning", "code": "LOAD_REPAIR", "nodeId": "x", "message": "m"}],
        {"sections": ALL, "runUncertainty": True, "topN": {"cutsets": 1, "importance": 1}},
    )
    assert json.dumps(core.get_data(), sort_keys=True) == before  # pure
    assert data["summary"]["headlineMethod"] == "mcub"
    assert data["cutsets"]["shown"] == 1 and data["cutsets"]["total"] == 2
    assert [m["id"] for m in data["importance"]["rows"]] == ["c"]  # sorted by FV
    assert data["uncertainty"]["mean"] == 1.5e-3
    assert fake_uncertainty.calls == [report_docx.MAX_REPORT_MC_N]
    assert [i["code"] for i in data["validation"]] == ["DANGLING_LINK", "LOAD_REPAIR"]
    assert fake_lint.calls[-1]["mode"] == "FTA"
    # The truncation signal in cutsets.truncation_signal's shape (the report's
    # own cut sets, expanded with the document's limits, are reused).
    assert fake_lint.calls[-1]["extra"]["cutsets"] == {
        "truncated": True, "truncatedBy": ["maxOrder"], "count": 2}
    assert data["traceability"][0]["requirementId"] == "REQ-1"
    assert data["assumptions"]["repeatedEvents"] == ["Valve"]
    assert {m["model"] for m in data["assumptions"]["models"]} == {"fixed", "rate"}
    assert data["unavailable"] == {}


def test_collect_without_the_other_workstreams(monkeypatch):
    import fta_web.cutsets
    import fta_web.lint

    monkeypatch.delattr(fta_web.cutsets, "compute", raising=False)
    monkeypatch.delattr(fta_web.lint, "run", raising=False)
    data = report_docx.collect_report_data(make_core(), [], {"sections": ALL})
    assert data["unavailable"]["cutsets"] == "unavailable"
    assert data["unavailable"]["importance"] == "unavailable"
    assert data["unavailable"]["uncertainty"] == "notRun"
    assert data["summary"]["headline"] is not None  # the engine stub works
    assert isinstance(data["validation"], list)  # falls back to session warnings


def test_uncertainty_n_is_capped(fakes):
    report_docx.collect_report_data(make_core(), [], {
        "sections": ["uncertainty"], "runUncertainty": True, "uncertaintyN": 10 ** 6})
    assert fake_uncertainty.calls == [report_docx.MAX_REPORT_MC_N]


# ---- the document ------------------------------------------------------------------------


def build(options, core=None, warnings=()):
    docx = pytest.importorskip("docx")
    data = report_docx.collect_report_data(core or make_core(), list(warnings), options)
    blob = report_docx.build_report(data, report_docx.normalize_options(options))
    assert blob[:2] == b"PK"
    return docx.Document(io.BytesIO(blob))


def test_docx_contains_headline_and_tables(fakes):
    doc = build({"sections": ALL, "runUncertainty": True, "sigFigs": 3})
    text = text_of(doc)
    assert "Fault Tree Analysis Report" in text
    assert "Plant" in text
    assert "0.0131" in text and "MCUB" in text  # headline + badge
    assert "Tree walk value: 0.0124" in text     # alt value
    assert "Mission time: 8760" in text
    assert "1 - exp(-λT)" in text
    assert "Priority-AND" in text and "non-coherent" in text
    assert "Sensor" in text and "60.0%" in text  # cut-set table
    assert "DANGLING_LINK" in text
    assert "REQ-1" in text
    assert "2.00e-7" in text  # event table, sig figs
    assert "Diagram not available." in text
    assert len(doc.tables) >= 6
    first_row = doc.tables[0].rows[0]._tr
    assert first_row.xpath("./w:trPr/w:tblHeader")  # header row repeats


def test_docx_event_table_has_the_source_column(fakes):
    doc = build({"sections": ["events"], "sigFigs": 3})
    table = doc.tables[-1]
    header = [c.text for c in table.rows[0].cells]
    assert header[-1] == "Source"
    pump = [r for r in table.rows if r.cells[0].text == "a"][0]
    assert pump.cells[len(header) - 1].text == "OREDA"


def test_docx_notes_unavailable_sections(monkeypatch):
    import fta_web.cutsets

    monkeypatch.delattr(fta_web.cutsets, "compute", raising=False)
    text = text_of(build({"sections": ["cutsets", "uncertainty"]}))
    assert "not available in this build" in text
    assert "Monte Carlo was not run" in text


def test_docx_japanese(fakes):
    from docx.oxml.ns import qn

    doc = build({"sections": ALL, "lang": "ja"})
    text = text_of(doc)
    assert "故障の木解析レポート" in text
    assert "ミニマルカットセット" in text and "前提条件" in text
    fonts = doc.styles["Normal"].element.rPr.find(qn("w:rFonts"))
    assert fonts.get(qn("w:eastAsia")) == "Meiryo"
    heading = doc.styles["Heading 1"].element.rPr.find(qn("w:rFonts"))
    assert heading.get(qn("w:eastAsia")) == "Meiryo"
    assert heading.get(qn("w:eastAsiaTheme")) is None


def test_docx_embeds_the_diagram(fakes):
    doc = build({"sections": ["diagram"], "diagramPng": tiny_png()})
    assert len(doc.inline_shapes) == 1
    assert "Diagram not available." not in text_of(doc)


def test_docx_in_eta_mode_skips_analysis(fakes):
    doc = build({"sections": ALL}, core=make_core(mode="ETA"))
    text = text_of(doc)
    assert "Event tree (ETA) mode" in text
    assert "Sensor | 1" not in text


def test_docx_uncertainty_shows_samples_and_a_partial_run(fakes):
    """A time-capped run is flagged, and the sample count is shown."""
    fake_uncertainty.completed = 1200
    text = text_of(build({"sections": ["uncertainty"], "runUncertainty": True,
                          "uncertaintyN": 5000}))
    assert "Time limit reached" in text
    assert "Samples | 1200 / 5000" in text
    assert "Seed | 42" in text


def test_docx_uncertainty_complete_run_is_not_flagged(fakes):
    text = text_of(build({"sections": ["uncertainty"], "runUncertainty": True,
                          "uncertaintyN": 300}))
    assert "Time limit reached" not in text
    assert "Samples | 300" in text


def test_report_time_limit_is_passed_and_capped(fakes):
    report_docx.collect_report_data(make_core(), [], {
        "sections": ["uncertainty"], "runUncertainty": True, "uncertaintyTimeLimit": 5})
    report_docx.collect_report_data(make_core(), [], {
        "sections": ["uncertainty"], "runUncertainty": True, "uncertaintyTimeLimit": 10 ** 6})
    report_docx.collect_report_data(make_core(), [], {
        "sections": ["uncertainty"], "runUncertainty": True, "uncertaintyTimeLimit": -1})
    assert fake_uncertainty.time_limits == [5.0, report_docx.MAX_REPORT_MC_SECONDS,
                                            report_docx.DEFAULT_REPORT_MC_SECONDS]


def test_docx_cutsets_show_mcub_and_rare_event(fakes):
    text = text_of(build({"sections": ["cutsets"], "sigFigs": 3}))
    assert "MCUB (all cut sets): 0.0124" in text
    assert "rare-event approximation: 0.0125" in text


def test_docx_infinite_rrw_is_shown_as_infinity(fakes, monkeypatch):
    import fta_web.importance

    def infinite(result):
        return [{"id": "c", "name": "Sensor", "q": 1.0, "fv": 1.0, "birnbaum": 1.0, "raw": 1.0,
                 "rrw": None, "rrwInfinite": True, "cutSetCount": 1}]

    monkeypatch.setattr(fta_web.importance, "compute", infinite)
    rows = [[c.text for c in r.cells] for r in build({"sections": ["importance"]}).tables[-1].rows]
    assert rows[1][6] == "∞"


def test_real_uncertainty_result_has_the_keys_the_report_reads():
    """Contract with uncertainty.run: the fakes above must not drift from it."""
    from fta_web import uncertainty

    tree = {"id": "root", "name": "t", "logicGate": "OR", "links": [], "children": [
        {"id": "a", "name": "A", "probability": 1e-3, "children": [], "links": [],
         "quant": {"model": "fixed", "unc": {"dist": "lognormal", "median": 1e-3, "ef": 3}}}]}
    result = uncertainty.run(tree, None, n=50, seed=1)
    for key in ("requested", "completed", "truncatedByTime", "seed", "method",
                "pointEstimate", "mean", "median", "p05", "p95", "std", "histogram"):
        assert key in result, key
    assert result["completed"] == 50 and result["truncatedByTime"] is False


def _truncating_core(max_order):
    """OR(AND(a, b), c): under maxOrder 1 the cut set {a, b} is dropped."""
    from fta_web.engine import WebCore

    core = WebCore()
    core.set_data({"id": "root", "name": "t", "type": "Root", "logicGate": "OR", "links": [],
                   "children": [
                       {"id": "g", "name": "G", "logicGate": "AND", "links": [], "children": [
                           {"id": "a", "name": "A", "probability": 0.1, "children": [], "links": []},
                           {"id": "b", "name": "B", "probability": 0.2, "children": [], "links": []}]},
                       {"id": "c", "name": "C", "probability": 0.3, "children": [], "links": []}]})
    core.set_analysis({"cutsets": {"maxOrder": max_order}})
    core.recalculate_probabilities()
    return core


def _validation_tab_issues(core):
    """What GET /api/analysis/validate and `cli validate` report."""
    from fta_web import cutsets, lint

    signal = cutsets.truncation_signal(core.get_data(), core.analysis)
    return lint.run(core.get_data(), core.analysis, [], extra={"cutsets": signal})


def _truncated(issues):
    return [i for i in issues if i["code"] == "CUTSETS_TRUNCATED"]


@pytest.mark.parametrize("sections", [["validation"], ["cutsets", "validation"]])
def test_report_validation_reports_truncation_like_the_validation_tab(sections):
    """B2B: the report's validation section and the Validation tab must list
    the same CUTSETS_TRUNCATED -- with the count of cut sets in its params
    (the report passed the whole cut-set result, so lint found no "count"
    and reported 0), and even when the cut-set section is not selected
    (then the report had no cut sets at all and never reported it)."""
    core = _truncating_core(1)
    tab = _truncated(_validation_tab_issues(core))
    assert tab and tab[0]["params"] == {"reason": "order", "count": 1}
    data = report_docx.collect_report_data(core, [], {"sections": sections})
    assert _truncated(data["validation"]) == tab


def test_report_validation_uses_the_document_limits_not_an_override():
    """USER_GUIDE: CUTSETS_TRUNCATED is about the cut sets "expanded with the
    document's limits". A report run with the Cut Sets tab's override limits
    must not change it (the cut-set section shows the override's own note)."""
    doc_truncates = _truncating_core(1)
    data = report_docx.collect_report_data(
        doc_truncates, [], {"sections": ["cutsets", "validation"], "limits": {"maxOrder": 6}})
    assert data["cutsets"]["truncated"] is False
    assert _truncated(data["validation"]) == _truncated(_validation_tab_issues(doc_truncates))

    override_truncates = _truncating_core(6)
    data = report_docx.collect_report_data(
        override_truncates, [], {"sections": ["cutsets", "validation"], "limits": {"maxOrder": 1}})
    assert data["cutsets"]["truncated"] is True
    assert _truncated(data["validation"]) == []
    assert _truncated(_validation_tab_issues(override_truncates)) == []


def test_real_importance_result_has_rrw_infinite():
    from fta_web import cutsets, importance

    tree = {"id": "root", "name": "t", "logicGate": "OR", "links": [], "children": [
        {"id": "a", "name": "A", "probability": 1e-3, "children": [], "links": []}]}
    rows = importance.compute(cutsets.compute(tree, None))
    assert rows[0]["rrwInfinite"] is True and rows[0]["rrw"] is None


# ---- the route -------------------------------------------------------------------------------


@pytest.fixture
def client(tmp_path):
    import flask

    from fta_web.routes.files import files_bp
    from fta_web.routes.report import report_bp
    from fta_web.routes.tree import tree_bp

    reset_state()
    app = flask.Flask(__name__)
    for bp in (tree_bp, files_bp, report_bp):
        app.register_blueprint(bp)
    app.config.update(TESTING=True)
    get_state().fs_root = tmp_path
    source = tmp_path / "My Plant.json"
    source.write_text(json.dumps({"title": "Doc", "date": "", "mode": "FTA",
                                  "tree": sample_tree()}), encoding="utf-8")
    with app.test_client() as test_client:
        assert test_client.post("/api/file/open", json={"path": str(source)}).status_code == 200
        yield test_client


def test_route_returns_an_attachment(client, fakes):
    docx = pytest.importorskip("docx")
    png = base64.b64encode(tiny_png()).decode("ascii")
    response = client.post("/api/report/docx", json={
        "sections": ALL, "sigFigs": 4, "lang": "en", "diagramPng": png,
        "topN": {"cutsets": 5}})
    assert response.status_code == 200, response.get_data(as_text=True)[:300]
    assert response.mimetype == report_docx.DOCX_MIME
    disposition = response.headers["Content-Disposition"]
    assert "attachment" in disposition and "My Plant_report.docx" in disposition
    doc = docx.Document(io.BytesIO(response.get_data()))
    assert len(doc.inline_shapes) == 1
    assert "0.01310" in text_of(doc)  # 4 sig figs


def test_route_503_without_python_docx(client, monkeypatch):
    from fta_web.routes import report

    monkeypatch.setattr(report, "docx_available", lambda: False)
    response = client.post("/api/report/docx", json={})
    assert response.status_code == 503
    error = response.get_json()["error"]
    assert error["code"] == "EXPORT_UNAVAILABLE"
    assert error["detail"] == {"format": "docx", "package": "python-docx",
                               "install": "uv sync --extra report"}


@pytest.mark.parametrize("body,field", [
    ({"diagramPng": base64.b64encode(b"GIF89a....").decode()}, "diagramPng"),
    ({"diagramPng": "!!!not base64!!!"}, "diagramPng"),
    ({"sections": ["nope"]}, "sections"),
    ({"sections": "metadata"}, "sections"),
    ({"lang": "fr"}, "lang"),
])
def test_route_rejects_bad_input(client, body, field):
    pytest.importorskip("docx")
    response = client.post("/api/report/docx", json=body)
    assert response.status_code == 400
    error = response.get_json()["error"]
    assert error["code"] == "INVALID_FIELD" and error["detail"]["field"] == field


def test_route_eta_is_allowed(client, fakes):
    docx = pytest.importorskip("docx")
    get_state().core.set_metadata(mode="ETA")
    response = client.post("/api/report/docx", json={"sections": ["headline", "cutsets"]})
    assert response.status_code == 200
    assert "Event tree (ETA) mode" in text_of(docx.Document(io.BytesIO(response.get_data())))


def test_route_renders_natively_when_no_png(client, fakes, monkeypatch):
    docx = pytest.importorskip("docx")
    monkeypatch.setattr(report_docx, "render_png", lambda dot: tiny_png() if dot else None)
    response = client.post("/api/report/docx", json={"sections": ["diagram"]})
    assert response.status_code == 200
    assert len(docx.Document(io.BytesIO(response.get_data())).inline_shapes) == 1
