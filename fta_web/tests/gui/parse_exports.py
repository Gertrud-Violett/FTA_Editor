"""Dump the numbers from a downloaded report .docx and Excel .xlsx as JSON.

Used by s_files.js (lib.parseExports) to compare the exports with the API:

    uv run --frozen --extra all --extra test python fta_web/tests/gui/parse_exports.py <report.docx> <export.xlsx>
"""
import json
import sys

import docx
import openpyxl


def parse_docx(path):
    d = docx.Document(path)
    out = {"headline": None, "tables": [], "headings": []}
    body = d.element.body
    # walk paragraphs and tables in document order
    current = None
    for child in body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            p = docx.text.paragraph.Paragraph(child, d)
            style = p.style.name if p.style is not None else ""
            if style.startswith("Heading"):
                current = p.text
                out["headings"].append(p.text)
            elif current in ("Top-event probability", "頂上事象の発生確率") and out["headline"] is None and p.runs:
                bold = [r.text for r in p.runs if r.bold]
                out["headline"] = {"value": bold[0] if bold else p.text, "text": p.text}
        elif tag == "tbl":
            t = docx.table.Table(child, d)
            rows = [[c.text for c in r.cells] for r in t.rows]
            out["tables"].append({"heading": current, "rows": rows})
    return out


def parse_xlsx(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    out = {"sheets": wb.sheetnames}
    if "Events" in wb.sheetnames:
        ws = wb["Events"]
        rows = list(ws.iter_rows(values_only=True))
        head = list(rows[0])
        out["events"] = [dict(zip(head, r)) for r in rows[1:]]
        # number formats of the probability columns
        ci = head.index("Calculated probability") + 1
        out["calcFormat"] = ws.cell(row=2, column=ci).number_format if ws.max_row >= 2 else None
    if "Analysis" in wb.sheetnames:
        ws = wb["Analysis"]
        out["analysis"] = {r[0]: r[1] for r in ws.iter_rows(min_row=2, values_only=True) if r and r[0]}
    return out


if __name__ == "__main__":
    print(json.dumps({"docx": parse_docx(sys.argv[1]), "xlsx": parse_xlsx(sys.argv[2])}, default=str, ensure_ascii=False))
